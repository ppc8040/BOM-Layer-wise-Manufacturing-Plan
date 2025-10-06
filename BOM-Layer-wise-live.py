import requests
import pandas as pd
import xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Border, Side, Alignment, PatternFill, Font
import logging
import warnings
import time
import os
import sys
from openpyxl import load_workbook
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
import msal
from datetime import datetime
from zoneinfo import ZoneInfo

if sys.platform == "win32":
    import os
    os.environ["PYTHONIOENCODING"] = "utf-8"

warnings.filterwarnings('ignore')

# --- Logging Setup ---
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s %(levelname)s %(message)s',
    handlers=[
        logging.FileHandler('bom_layerwise.log'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# --- Configuration ---
SAP_API_URL = "https://my409486-api.s4hana.cloud.sap/sap/opu/odata/sap/API_BILL_OF_MATERIAL_SRV;v=0002"
USERNAME = "INDUS_OPEN_ORDER"
PASSWORD = "w6fo3zfq/bDYr}8>BU@{i3oFd[dq#Zv\\/eL&}SKF"
VERIFY_SSL = True
PREFERRED_FORMAT = 'json'
BATCH_SIZE = 500
MAX_LAYERS = 4

# === SHAREPOINT-ONLY GRAPH API INTEGRATION ===
import msal
import json

class AuthConfig:
    CLIENT_ID = 'a2239be7-12cd-442d-983f-ea7e316ef767'
    CLIENT_SECRET = 'sho8Q~sanlYEX3SKXjYJb0NSXL.gPRpdi63-tcmp'
    AUTHORITY = 'https://login.microsoftonline.com/ed97e9bb-e119-4bd4-ab00-307d64bdf908'

def get_access_token():
    """Get Microsoft Graph API access token"""
    auth_config = AuthConfig()
    app = msal.ConfidentialClientApplication(
        auth_config.CLIENT_ID, 
        authority=auth_config.AUTHORITY, 
        client_credential=auth_config.CLIENT_SECRET
    )
    
    result = app.acquire_token_silent(scopes=['https://graph.microsoft.com/.default'], account=None)
    if not result:
        result = app.acquire_token_for_client(scopes=['https://graph.microsoft.com/.default'])
    
    if 'access_token' in result:
        return result['access_token']
    else:
        raise Exception("Failed to acquire access token.")

def get_excel_column_letter(col_num):
    """Convert column number to Excel column letter (e.g., 1 -> A, 27 -> AA)"""
    letter = ""
    while col_num > 0:
        col_num -= 1
        letter = chr(col_num % 26 + ord('A')) + letter
        col_num //= 26
    return letter

def create_sharepoint_worksheet_optimized(site_id, file_server_relative_url, worksheet_name, df, access_token, session_id, chunk_size=1000, max_retries=5):
    """Robust creation/update for large SharePoint sheets with data validation and exponential backoff"""
    session = requests.Session()
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": session_id
    }

    logger.info(f"Creating/updating SharePoint worksheet: {worksheet_name}")
    # Check for worksheet existence, create if needed (as before)
    ws_list_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets"
    ws_list_resp = session.get(ws_list_url, headers=api_headers)
    ws_names = [ws['name'] for ws in ws_list_resp.json().get('value', [])] if ws_list_resp.status_code == 200 else []
    if worksheet_name not in ws_names:
        create_ws_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets/add"
        session.post(create_ws_url, headers=api_headers, json={"name": worksheet_name})
        time.sleep(2)
    # Clear data in full
    try:
        clear_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/usedRange/clear"
        session.post(clear_url, headers=api_headers, json={"applyTo": "Contents"}, timeout=60)
        time.sleep(1)
    except Exception as clear_error:
        logger.warning(f"Clear operation failed: {clear_error}")

    # Write headers first
    headers = df.columns.tolist()
    data_conv = df.astype(str).where(~df.isna(), "")  # Ensures all strings
    num_cols = len(headers)
    header_range = f"A1:{get_excel_column_letter(num_cols)}1"
    header_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{header_range}')"
    session.patch(header_url, headers=api_headers, json={"values": [headers]}, timeout=60)

    # Chunked rows
    total_rows = len(df)
    total_chunks = (total_rows + chunk_size - 1) // chunk_size
    for chunk_idx in range(0, total_rows, chunk_size):
        chunk_end = min(chunk_idx + chunk_size, total_rows)
        chunk_data = data_conv.iloc[chunk_idx:chunk_end].values.tolist()
        chunk_range = f"A{chunk_idx+2}:{get_excel_column_letter(num_cols)}{chunk_end+1}"
        chunk_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{chunk_range}')"
        # Exponential backoff
        for attempt in range(max_retries):
            try:
                chunk_response = session.patch(
                    chunk_url, headers=api_headers, json={"values": chunk_data}, timeout=120
                )
                if chunk_response.status_code in [200, 201]:
                    logger.info(f"Chunk {(chunk_idx//chunk_size)+1}/{total_chunks} successful for '{worksheet_name}' (rows {chunk_idx+2}-{chunk_end+1})")
                    break
                elif chunk_response.status_code in {429, 503, 504}:
                    wait = 2 ** attempt
                    logger.warning(f"Throttling or timeout on chunk {(chunk_idx//chunk_size)+1}, wait {wait}s")
                    time.sleep(wait)
                else:
                    if attempt == max_retries - 1:
                        logger.error(f"Chunk {(chunk_idx//chunk_size)+1} for '{worksheet_name}' failed after {max_retries} attempts: {chunk_response.text}")
                    else:
                        time.sleep(2)
            except Exception as ex:
                wait = 2 ** attempt
                logger.warning(f"Exception on chunk {(chunk_idx//chunk_size)+1}, attempt {attempt+1}: {ex}, waiting {wait}s")
                time.sleep(wait)
        else:
            logger.error(f"Giving up on chunk {(chunk_idx//chunk_size)+1} for '{worksheet_name}'")
            return False
    logger.info(f"'{worksheet_name}' data upload complete.")
    apply_sharepoint_formatting_fixed(site_id, file_server_relative_url, worksheet_name, session_id, access_token)
    return True
def apply_sharepoint_formatting_fixed(site_id, file_server_relative_url, worksheet_name, session_id, access_token):
    """Apply formatting to SharePoint worksheet with proper border coverage"""
    session = requests.Session()
    
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": session_id
    }
    
    try:
        # Get the actual used range from the server
        used_range_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/usedRange"
        resp = session.get(used_range_url, headers=api_headers)
        if resp.status_code != 200:
            logger.warning(f"Could not retrieve usedRange for '{worksheet_name}': {resp.status_code}")
            return
        
        address = resp.json().get("address", "")
        if "!" in address:
            address = address.split("!", 1)[1]
        
        # Apply right alignment for all data
        align_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{address}')/format"
        align_payload = {"horizontalAlignment": "Right"}
        session.patch(align_url, headers=api_headers, json=align_payload)
        
        # Apply borders to entire used range
        border_ids = ["EdgeTop", "EdgeBottom", "EdgeLeft", "EdgeRight", "InsideHorizontal", "InsideVertical"]
        for border_id in border_ids:
            border_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{address}')/format/borders('{border_id}')"
            border_payload = {"style": "Continuous", "color": "#000000", "weight": "Thin"}
            session.patch(border_url, headers=api_headers, json=border_payload)
        
        # Get header range
        if address:
            cols = address.split(':')
            if len(cols) == 2:
                start_col = cols[0][0]  # Get first column letter
                end_col = cols[1][0]    # Get  column letter
                header_range = f"{start_col}1:{end_col}1"
                
                # Format header row
                header_fill_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{header_range}')/format/fill"
                header_fill_payload = {"color": "#366092"}
                session.patch(header_fill_url, headers=api_headers, json=header_fill_payload)
                
                header_font_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{header_range}')/format/font"
                header_font_payload = {"bold": True, "color": "#FFFFFF"}
                session.patch(header_font_url, headers=api_headers, json=header_font_payload)
        
        # Auto-fit columns
        autofit_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{address}')/format/autofitColumns"
        session.post(autofit_url, headers=api_headers)
        
        logger.info(f"Applied formatting with full border coverage to '{worksheet_name}' (range: {address})")
        
    except Exception as e:
        logger.warning(f"Error applying formatting to '{worksheet_name}': {e}")

def create_sharepoint_bom_layerwise_sheet(site_id, file_server_relative_url, session_id, access_token, df_bom, all_paths, max_layers, header_mapping, header_qty_mapping):
    """Create BOM Layer-wise sheet directly in SharePoint"""
    logger.info("Creating BOM Layer-wise sheet directly in SharePoint...")
    
    # Same processing logic as local export_tree_paths
    columns = ["Main_FG", "Main_FG_Description"]
    for lvl in range(1, max_layers+1):
        columns.extend([
            f"BOM_Layer_{lvl}",
            f"BOM_Layer_{lvl}_Desc", 
            f"BOM_Layer_{lvl}_Quantity",
            f"BOM_Layer_{lvl}_Unit"
        ])
    
    records = []
    for path in all_paths:
        row = dict()
        if path:
            first = path[0]
            row["Main_FG"] = first[0]
            row["Main_FG_Description"] = header_mapping.get(first[0], "")
        else:
            row["Main_FG"], row["Main_FG_Description"] = "", ""
        
        for i, comp in enumerate(path):
            if i < max_layers:
                row[f"BOM_Layer_{i+1}"] = comp[4]
                row[f"BOM_Layer_{i+1}_Desc"] = comp[1]
                row[f"BOM_Layer_{i+1}_Quantity"] = comp[2]
                row[f"BOM_Layer_{i+1}_Unit"] = comp[3]
        
        records.append(row)
    
    df_out = pd.DataFrame(records, columns=columns)
    
    # Apply same filtering as local
    filtered_df_out = df_out[df_out["Main_FG"].str.startswith(("A", "C", "I"))].copy()
    logger.info(f"Filtered rows (A/C/I only), remaining count: {len(filtered_df_out)}")
    
    # Apply BOM header quantity division
    def divide_qtys_by_divisor(qty_str, divisor):
        try:
            if not isinstance(qty_str, str):
                qty_str = str(qty_str) if qty_str is not None else ''
            qtys = [float(q) if q.strip() else 0 for q in qty_str.split(" | ")]
            divided = [q / divisor if divisor != 0 else q for q in qtys]
            result = " | ".join([str(round(q, 10)).rstrip('0').rstrip('.') if q != 0 else '0' for q in divided])
            return result
        except Exception as e:
            return qty_str
    
    for main_fg, group in filtered_df_out.groupby("Main_FG"):
        divisor = header_qty_mapping.get(main_fg, 1)
        try:
            divisor = float(divisor)
        except Exception:
            divisor = 1
        
        if pd.notna(divisor) and divisor != 0:
            mask = filtered_df_out["Main_FG"] == main_fg
            filtered_df_out.loc[mask, "BOM_Layer_1_Quantity"] = filtered_df_out.loc[mask, "BOM_Layer_1_Quantity"].apply(divide_qtys_by_divisor, divisor=divisor)
    
    # Convert G to KG for all layers
    for lvl in range(1, max_layers + 1):
        qty_col = f"BOM_Layer_{lvl}_Quantity"
        unit_col = f"BOM_Layer_{lvl}_Unit"
        
        if qty_col in filtered_df_out.columns and unit_col in filtered_df_out.columns:
            mask = filtered_df_out[unit_col].astype(str).str.strip().str.upper() == 'G'
            n_g = mask.sum()
            
            if n_g > 0:
                def convert_and_scale_g_to_kg(row):
                    qty_value = row[qty_col]
                    try:
                        q = float(qty_value)
                    except:
                        q = 0.0
                    return q / 1000
                
                filtered_df_out.loc[mask, qty_col] = filtered_df_out.loc[mask].apply(convert_and_scale_g_to_kg, axis=1)
                filtered_df_out.loc[mask, unit_col] = 'KG'
    
    # Add timestamp column
    filtered_df_out['Last_Updated'] = datetime.now(ZoneInfo("Asia/Dubai")).strftime('%Y-%m-%d %H:%M:%S')
    
    return create_sharepoint_worksheet_optimized(site_id, file_server_relative_url, "BOM Layer-wise", filtered_df_out, access_token, session_id)

def create_sharepoint_fg_rm_sheet(site_id, file_server_relative_url, session_id, access_token, main_sheet_df, max_layers):
    """Create FG-RM Required sheet directly in SharePoint"""
    logger.info("Creating FG-RM Required sheet directly in SharePoint...")
    
    # Same logic as local export_fg_rm_summary
    summary_records = []
    
    for main_fg in main_sheet_df["Main_FG"].unique():
        if not isinstance(main_fg, str) or not main_fg:
            continue
        
        fg_rows = main_sheet_df[main_sheet_df["Main_FG"] == main_fg]
        rm_dict = {}
        
        for _, row in fg_rows.iterrows():
            layer_1_qty_str = str(row.get("BOM_Layer_1_Quantity", "1")).strip()
            
            try:
                layer_1_qtys = [float(q) if q.strip() else 1 for q in layer_1_qty_str.split(" | ")]
                layer_1_total_qty = sum(layer_1_qtys) if layer_1_qtys else 1
            except Exception:
                layer_1_total_qty = 1
            
            for lvl in range(1, max_layers + 1):
                mat_col = f"BOM_Layer_{lvl}"
                desc_col = f"BOM_Layer_{lvl}_Desc"
                qty_col = f"BOM_Layer_{lvl}_Quantity"
                unit_col = f"BOM_Layer_{lvl}_Unit"
                
                mat = str(row.get(mat_col, '')).strip()
                if mat.startswith('5'):
                    desc = str(row.get(desc_col, '')).strip()
                    qty_str = str(row.get(qty_col, '')).strip()
                    unit = str(row.get(unit_col, '')).strip().upper()
                    
                    if 'pallet' in desc.lower() or 'box' in desc.lower() or 'label' in desc.lower() or 'paint' in desc.lower():
                        continue
                    
                    try:
                        qtys = [float(q) if q.strip() else 0 for q in qty_str.split(" | ")]
                        component_qty = sum(qtys)
                    except Exception:
                        component_qty = 0
                    
                    if component_qty > 0:
                        if lvl == 1:
                            final_qty = component_qty
                        else:
                            final_qty = component_qty * layer_1_total_qty
                        
                        if mat in rm_dict:
                            rm_dict[mat]['qty'] += final_qty
                        else:
                            rm_dict[mat] = {
                                'desc': desc,
                                'qty': final_qty,
                                'unit': unit
                            }
        
        for rm_code, details in rm_dict.items():
            if details['qty'] > 0:
                summary_records.append({
                    "Main_FG": main_fg,
                    "Raw Material": rm_code,
                    "Raw Material Description": details['desc'],
                    "Raw Material qty": details['qty'],
                    "Raw Material UoM": details['unit']
                })
    
    summary_df = pd.DataFrame(summary_records, columns=[
        "Main_FG", "Raw Material", "Raw Material Description", "Raw Material qty", "Raw Material UoM"
    ])
    
    # Add timestamp
    summary_df['Last_Updated'] = datetime.now(ZoneInfo("Asia/Dubai")).strftime('%Y-%m-%d %H:%M:%S')
    
    return create_sharepoint_worksheet_optimized(site_id, file_server_relative_url, "FG-RM Required", summary_df, access_token, session_id, chunk_size=1000)

def create_sharepoint_mb52_sheet(site_id, file_server_relative_url, session_id, access_token, stock_df):
    """Create MB52 Raw Data sheet directly in SharePoint"""
    logger.info("Creating MB52 Raw Data sheet directly in SharePoint...")
    
    # Add timestamp column
    stock_df_with_timestamp = stock_df.copy()
    stock_df_with_timestamp['Last_Updated'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    
    return create_sharepoint_worksheet_optimized(site_id, file_server_relative_url, "MB52 Raw Data", stock_df_with_timestamp, access_token, session_id)

def read_sharepoint_production_tracker(site_id, file_server_relative_url, session_id, access_token):
    """Read Production Tracker data robustly from SharePoint; return empty DF gracefully on error."""
    session = requests.Session()
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": session_id
    }
    range_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/usedRange"
    try:
        range_response = session.get(range_url, headers=api_headers, timeout=120)
        if range_response.status_code != 200:
            logger.error(f"Failed to read Production Tracker. Status: {range_response.status_code}. Will initialize as empty DataFrame.")
            return pd.DataFrame(), pd.DataFrame()
        range_data = range_response.json()
        values = range_data.get('values', [])
        if not values:
            logger.warning("No data found in Production Tracker")
            return pd.DataFrame(), pd.DataFrame()
        headers = values[0]
        data_rows = values[1:]
        df = pd.DataFrame(data_rows, columns=headers)
        df['Indus Part #'] = df['Indus Part #'].astype(str).str.strip().str.upper()
        if not pd.api.types.is_datetime64_any_dtype(df['DI']):
            df['DI'] = pd.to_datetime(df['DI'], errors='coerce')
        df['Status'] = df['Status'].astype(object)
        def is_blank(val):
            if pd.isna(val): return True
            if isinstance(val, str) and val.strip() == "": return True
            if isinstance(val, str) and val.strip().lower() in ("done", "cancelled"):
                return False
            return False
        mask_blank = df['Status'].apply(is_blank)
        mask_valid = (
            df['Indus Part #'].notna() &
            (df['Indus Part #'] != '') &
            (pd.to_numeric(df['Qty'], errors='coerce') > 0) &
            df['DI'].notna()
        )
        df_filtered = df[mask_blank & mask_valid].sort_values(by=['DI'], ascending=[True])
        logger.info(f"Read {len(df_filtered)} valid FIFO rows from SharePoint Production Tracker")
        return df_filtered, df
    except Exception as e:
        logger.error(f"Exception fetching Production Tracker: {e}")
        return pd.DataFrame(), pd.DataFrame()


def read_sharepoint_worksheet_data(site_id, file_server_relative_url, worksheet_name, session_id, access_token):
    """Generic function to read data from any SharePoint worksheet"""
    session = requests.Session()
    
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": session_id
    }
    
    try:
        # Get used range from specified worksheet
        range_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/usedRange"
        
        range_response = session.get(range_url, headers=api_headers, timeout=120)
        if range_response.status_code != 200:
            logger.error(f"Failed to read {worksheet_name}: {range_response.status_code}")
            return pd.DataFrame()
        
        range_data = range_response.json()
        values = range_data.get('values', [])
        
        if not values:
            logger.warning(f"No data found in {worksheet_name}")
            return pd.DataFrame()
        
        # Convert to DataFrame
        headers = values[0]
        data_rows = values[1:]
        df = pd.DataFrame(data_rows, columns=headers)
        
        logger.info(f"Read {len(df)} rows from SharePoint {worksheet_name}")
        return df
        
    except Exception as e:
        logger.error(f"Error reading SharePoint {worksheet_name}: {e}")
        return pd.DataFrame()

def update_sharepoint_production_tracker_results(site_id, file_server_relative_url, session_id, access_token, results_df, full_df):
    """Update Production Tracker columns P and Q with FIFO allocation results"""
    session = requests.Session()
    
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": session_id
    }
    
    try:
        # Set headers for columns P and Q
        header_updates = [
            {"range": "P1", "values": [["Material Available"]]},
            {"range": "Q1", "values": [["Shortage Details"]]}
        ]
        
        for update in header_updates:
            update_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{update['range']}')"
            session.patch(update_url, headers=api_headers, json=update, timeout=60)
        
        # Create results mapping
        results_mapping = {
            int(result['index']): {
                'status': result['Status'],
                'shortage': result['Shortage_Details']
            } for _, result in results_df.iterrows()
        }
        
        # Prepare data for ALL rows
        total_rows = len(full_df)
        p_values = []
        q_values = []
        
        for idx, row in full_df.iterrows():
            status_raw = row['Status']
            status_clean = str(status_raw).strip().lower() if isinstance(status_raw, str) else ""
            
            if status_clean in ("done", "cancelled"):
                p_values.append("NA")
                q_values.append("NA")
            elif idx in results_mapping:
                p_values.append(results_mapping[idx]['status'])
                q_values.append(results_mapping[idx]['shortage'])
            else:
                p_values.append("NA")
                q_values.append("NA")
        
        # Update entire columns P and Q at once
        if total_rows > 0:
            # Update column P
            p_range = f"P2:P{total_rows + 1}"
            p_values_formatted = [[val] for val in p_values]
            p_update_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{p_range}')"
            p_payload = {"values": p_values_formatted}
            
            p_response = session.patch(p_update_url, headers=api_headers, json=p_payload, timeout=300)
            if p_response.status_code in [200, 201]:
                logger.info("Column P (Material Available) updated successfully")
            
            # Update column Q
            q_range = f"Q2:Q{total_rows + 1}"
            q_values_formatted = [[val] for val in q_values]
            q_update_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{q_range}')"
            q_payload = {"values": q_values_formatted}
            
            q_response = session.patch(q_update_url, headers=api_headers, json=q_payload, timeout=300)
            if q_response.status_code in [200, 201]:
                logger.info("Column Q (Shortage Details) updated successfully")
        
        # Apply formatting to columns P and Q
        apply_production_tracker_formatting(site_id, file_server_relative_url, session_id, access_token, total_rows)
        
        updated_count = len([idx for idx in range(len(full_df)) if idx in results_mapping])
        logger.info(f"Production Tracker updated with {updated_count} allocation results out of {total_rows} total rows")
        return True
        
    except Exception as e:
        logger.error(f"Error updating SharePoint Production Tracker: {e}")
        return False

def apply_production_tracker_formatting(site_id, file_server_relative_url, session_id, access_token, total_rows):
    """Apply formatting to Production Tracker columns P and Q"""
    session = requests.Session()
    
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": session_id
    }
    
    try:
        # Format headers (P1:Q1)
        header_range = "P1:Q1"
        
        header_fill_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{header_range}')/format/fill"
        header_fill_payload = {"color": "#366092"}
        session.patch(header_fill_url, headers=api_headers, json=header_fill_payload)
        
        header_font_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{header_range}')/format/font"
        header_font_payload = {"bold": True, "color": "#FFFFFF"}
        session.patch(header_font_url, headers=api_headers, json=header_font_payload)
        
        # Data range formatting
        if total_rows > 0:
            data_range = f"P2:Q{total_rows + 1}"
            
            # Apply borders
            border_ids = ["EdgeTop", "EdgeBottom", "EdgeLeft", "EdgeRight", "InsideHorizontal", "InsideVertical"]
            for border_id in border_ids:
                border_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{data_range}')/format/borders('{border_id}')"
                border_payload = {"style": "Continuous", "color": "#000000", "weight": "Thin"}
                session.patch(border_url, headers=api_headers, json=border_payload)
            
            # Auto-fit columns P and Q
            autofit_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='P:Q')/format/autofitColumns"
            session.post(autofit_url, headers=api_headers)
        
        logger.info("Production Tracker columns P and Q formatted successfully")
        
    except Exception as e:
        logger.warning(f"Error formatting Production Tracker columns: {e}")

def set_sharepoint_timestamp(site_id, file_server_relative_url, worksheet_name, session_id, access_token, cell_address="R1"):
    """Set timestamp in specified cell"""
    session = requests.Session()
    
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": session_id
    }
    
    update_datetime = datetime.now(ZoneInfo("Asia/Dubai")).strftime('Last Updated: %d-%m-%Y %H:%M:%S')
    payload = {"values": [[update_datetime]]}
    
    update_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{cell_address}')"
    
    try:
        r = session.patch(update_url, headers=api_headers, json=payload, timeout=60)
        if r.status_code in [200, 201]:
            logger.info(f"Timestamp written to cell {cell_address} in '{worksheet_name}'")
            return True
        else:
            logger.error(f"Failed to update timestamp: {r.status_code}")
            return False
    except Exception as e:
        logger.error(f"Error setting timestamp: {e}")
        return False

# --- BOM Item Data Fetch (SAME AS ORIGINAL) ---
def fetch_bom_data(api_url, username, password, verify_ssl=True, preferred_format='json', batch_size=500, max_workers=15) -> pd.DataFrame:
    logger.info("Starting BOM ITEM data fetch with enhanced concurrency...")
    session = requests.Session()
    session.auth = (username, password)
    session.verify = verify_ssl
    session.headers.update({'Accept': f'application/{preferred_format}', 'Content-Type': f'application/{preferred_format}'})
    
    select_fields = [
        'Material', 'ComponentDescription', 'BillOfMaterialComponent',
        'IsAssembly', 'BillOfMaterialItemQuantity', 'BillOfMaterialItemUnit'
    ]
    
    endpoint = f"{api_url}/MaterialBOMItem"
    
    def fetch_batch(skip):
        params = {'$select': ','.join(select_fields), '$top': batch_size, '$orderby': 'Material', '$skip': skip}
        try:
            logger.info(f"Fetching BOM items batch starting at {skip}...")
            response = session.get(endpoint, params=params, timeout=180)
            response.raise_for_status()
            
            content_type = response.headers.get('content-type', '')
            if 'xml' in content_type or preferred_format == 'xml':
                return parse_xml_bom(response.content)
            else:
                return parse_json_bom(response.text)
        except Exception as e:
            logger.error(f"Failed to fetch batch at skip={skip} : {e}")
            return []
    
    # Estimate total items
    params_count = {'$top': 1}
    resp_count = session.get(endpoint, params=params_count, timeout=30)
    resp_count.raise_for_status()
    data_count = resp_count.json()
    total_items = int(data_count.get('d', {}).get('__count', '0')) if '__count' in data_count.get('d', {}) else 0
    if total_items == 0:
        total_items = 50000
    
    total_batches = (total_items // batch_size) + 1
    all_records = []
    
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {executor.submit(fetch_batch, batch_num * batch_size): batch_num for batch_num in range(total_batches)}
        for future in as_completed(futures):
            records = future.result()
            if not records:
                logger.warning(f"No records fetched for batch {futures[future]}")
            else:
                all_records.extend(records)
                logger.info(f"Fetched batch size: {len(records)} (total so far: {len(all_records)})")
    
    logger.info(f"Total BOM ITEM records fetched: {len(all_records)}")
    return pd.DataFrame(all_records)

def fetch_material_bom_headers(api_url, username, password, verify_ssl=True, preferred_format='json', batch_size=500):
    logger.info("Starting BOM HEADER data fetch...")
    session = requests.Session()
    session.auth = (username, password)
    session.verify = verify_ssl
    session.headers.update({'Accept': f'application/{preferred_format}', 'Content-Type': f'application/{preferred_format}'})
    
    select_fields = ['Material', 'ProductDescription', 'BOMHeaderQuantityInBaseUnit']
    endpoint = f"{api_url}/MaterialBOM"
    params = {'$select': ','.join(select_fields), '$top': batch_size}
    
    all_records = []
    skip = 0
    
    while True:
        params['$skip'] = skip
        logger.info(f"Fetching BOM headers batch starting at {skip}...")
        response = session.get(endpoint, params=params, timeout=180)
        response.raise_for_status()
        
        content_type = response.headers.get('content-type', '')
        if 'xml' in content_type or preferred_format == 'xml':
            ns = {'atom': 'http://www.w3.org/2005/Atom', 'd': 'http://schemas.microsoft.com/ado/2007/08/dataservices', 'm': 'http://schemas.microsoft.com/ado/2007/08/dataservices/metadata'}
            tree = ET.fromstring(response.content)
            entries = tree.findall('.//atom:entry', ns)
            records = []
            for entry in entries:
                properties = entry.find('.//m:properties', ns)
                if properties is not None:
                    m = properties.find('.//d:Material', ns)
                    d = properties.find('.//d:ProductDescription', ns)
                    h = properties.find('.//d:BOMHeaderQuantityInBaseUnit', ns)
                    rec = {
                        'Material': m.text.strip() if m is not None and m.text else '',
                        'ProductDescription': d.text.strip() if d is not None and d.text else '',
                        'BOMHeaderQuantityInBaseUnit': h.text.strip() if h is not None and h.text else ''
                    }
                    records.append(rec)
        else:
            import json
            dct = response.json()
            records = [{'Material': i.get('Material', ''), 'ProductDescription': i.get('ProductDescription', ''), 'BOMHeaderQuantityInBaseUnit': i.get('BOMHeaderQuantityInBaseUnit', '')}
                       for i in dct.get('d', {}).get('results', [])]
        
        if not records:
            logger.warning("No more BOM header records found.")
            break
        
        all_records.extend(records)
        logger.info(f"Fetched header batch size: {len(records)} (total so far: {len(all_records)})")
        
        if len(records) < batch_size:
            break
        skip += batch_size
    
    logger.info(f"Total BOM HEADER records fetched: {len(all_records)}")
    return pd.DataFrame(all_records)

def fetch_material_stock_data(api_url, username, password, verify_ssl=True, preferred_format='json', batch_size=500):
    logger.info("Starting Material Stock data fetch...")
    session = requests.Session()
    session.auth = (username, password)
    session.verify = verify_ssl
    session.headers.update({'Accept': f'application/{preferred_format}', 'Content-Type': f'application/{preferred_format}'})
    
    select_fields = [
        'Material', 'Plant', 'StorageLocation', 'Batch',
        'InventorySpecialStockType',
        'InventoryStockType', 'MaterialBaseUnit', 'MatlWrhsStkQtyInMatlBaseUnit'
    ]
    
    stock_api_url = "https://my409486-api.s4hana.cloud.sap/sap/opu/odata/sap/API_MATERIAL_STOCK_SRV"
    endpoint = f"{stock_api_url}/A_MatlStkInAcctMod"
    params = {'$select': ','.join(select_fields), '$top': batch_size}
    
    all_records = []
    skip = 0
    
    while True:
        params['$skip'] = skip
        logger.info(f"Fetching Material Stock batch starting at {skip}...")
        
        try:
            response = session.get(endpoint, params=params, timeout=180)
            response.raise_for_status()
            
            content_type = response.headers.get('content-type', '')
            if 'xml' in content_type or preferred_format == 'xml':
                records = parse_xml_material_stock(response.content)
            else:
                records = parse_json_material_stock(response.text)
            
            if not records:
                logger.warning("No more Material Stock records found.")
                break
            
            all_records.extend(records)
            logger.info(f"Fetched Material Stock batch size: {len(records)} (total so far: {len(all_records)})")
            
            if len(records) < batch_size:
                break
            skip += batch_size
        except Exception as e:
            logger.error(f"Failed to fetch Material Stock batch at skip={skip}: {e}")
            break
    
    logger.info(f"Total Material Stock records fetched: {len(all_records)}")
    return pd.DataFrame(all_records)

def parse_json_bom(text):
    import json
    try:
        data = json.loads(text)
        results = []
        for item in data.get('d', {}).get('results', []):
            r = {}
            for key in ['Material', 'BillOfMaterialComponent', 'ComponentDescription', 'IsAssembly', 'BillOfMaterialItemQuantity', 'BillOfMaterialItemUnit']:
                value = item.get(key, '')
                if isinstance(value, dict): value = str(value)
                r[key] = value
            results.append(r)
        return results
    except Exception as ex:
        logger.error(f"JSON BOM parse failed: {str(ex)}")
        return []

def parse_xml_bom(xml_content):
    ns = {'atom': 'http://www.w3.org/2005/Atom', 'd': 'http://schemas.microsoft.com/ado/2007/08/dataservices', 'm': 'http://schemas.microsoft.com/ado/2007/08/dataservices/metadata'}
    tree = ET.fromstring(xml_content)
    records = []
    entries = tree.findall('.//atom:entry', ns)
    
    for entry in entries:
        properties = entry.find('.//m:properties', ns)
        if properties is not None:
            record = {}
            for field in ['Material', 'BillOfMaterialComponent', 'ComponentDescription', 'IsAssembly', 'BillOfMaterialItemQuantity', 'BillOfMaterialItemUnit']:
                elem = properties.find(f'.//d:{field}', ns)
                record[field] = elem.text.strip() if elem is not None and elem.text else ''
            records.append(record)
    return records

def parse_json_material_stock(text):
    import json
    try:
        data = json.loads(text)
        results = []
        for item in data.get('d', {}).get('results', []):
            r = {}
            for key in ['Material', 'Plant', 'StorageLocation', 'Batch', 'InventorySpecialStockType', 'InventoryStockType', 'MaterialBaseUnit', 'MatlWrhsStkQtyInMatlBaseUnit']:
                value = item.get(key, '')
                if isinstance(value, dict):
                    value = str(value)
                r[key] = value
            results.append(r)
        return results
    except Exception as ex:
        logger.error(f"JSON Material Stock parse failed: {str(ex)}")
        return []

def parse_xml_material_stock(xml_content):
    ns = {'atom': 'http://www.w3.org/2005/Atom', 'd': 'http://schemas.microsoft.com/ado/2007/08/dataservices', 'm': 'http://schemas.microsoft.com/ado/2007/08/dataservices/metadata'}
    tree = ET.fromstring(xml_content)
    records = []
    entries = tree.findall('.//atom:entry', ns)
    
    for entry in entries:
        properties = entry.find('.//m:properties', ns)
        if properties is not None:
            record = {}
            for field in ['Material', 'Plant', 'StorageLocation', 'Batch', 'InventorySpecialStockType', 'InventoryStockType', 'MaterialBaseUnit', 'MatlWrhsStkQtyInMatlBaseUnit']:
                elem = properties.find(f'.//d:{field}', ns)
                record[field] = elem.text.strip() if elem is not None and elem.text else ''
            records.append(record)
    return records

def standardize_bom_data(df):
    logger.info("Standardizing BOM item dataframe...")
    df = df.fillna('')
    
    for col in ['Material', 'BillOfMaterialComponent', 'ComponentDescription', 'BillOfMaterialItemUnit']:
        df[col] = df[col].astype(str).str.strip().str.upper()
    
    df['BillOfMaterialItemQuantity'] = pd.to_numeric(df['BillOfMaterialItemQuantity'], errors='coerce').fillna(0)
    df['IsAssembly'] = df['IsAssembly'].astype(str).str.upper()
    df = df.drop_duplicates(subset=['Material', 'BillOfMaterialComponent'], keep='first')
    
    logger.info(f"BOM item dataframe standardized, total rows: {len(df)}")
    return df

def tree_explode(df, start_mat, max_depth, parents_seen=None, path=None):
    if parents_seen is None:
        parents_seen = set()
    if path is None:
        path = []
    
    children = df[df["Material"] == start_mat]
    
    if len(children) == 0:
        yield path
    else:
        for idx, row in children.iterrows():
            child = row["BillOfMaterialComponent"]
            if child == "" or child in parents_seen:
                continue
            
            row_tuple = (
                start_mat,
                row["ComponentDescription"],
                row["BillOfMaterialItemQuantity"],
                row["BillOfMaterialItemUnit"],
                child,
                row["IsAssembly"]
            )
            
            next_path = path + [row_tuple]
            
            if row["IsAssembly"] == 'X' and len(next_path) < max_depth:
                yield from tree_explode(df, child, max_depth, parents_seen | {start_mat}, next_path)
            else:
                yield next_path

def create_mb52_stock_dict(mb52_df):
    """Create stock dictionary from MB52 data for Plant 1000 ONLY"""
    logger.info("Creating MB52 stock dictionary for Plant 1000...")
    
    plant_1000_stock = mb52_df[
        mb52_df['Plant'].astype(str).str.strip() == '1000'
    ].copy()
    
    logger.info(f"Plant 1000 stock records: {len(plant_1000_stock)}")
    
    qty_col = 'MatlWrhsStkQtyInMatlBaseUnit'
    if qty_col in plant_1000_stock.columns:
        plant_1000_stock[qty_col] = pd.to_numeric(plant_1000_stock[qty_col], errors='coerce').fillna(0)
    else:
        logger.error(f"Column {qty_col} not found in stock data")
        return {}
    
    # Aggregate by Material
    stock_dict = {}
    for _, row in plant_1000_stock.iterrows():
        material = str(row['Material']).strip().upper()
        qty = float(row[qty_col])
        
        if material in stock_dict:
            stock_dict[material] += qty
        else:
            stock_dict[material] = qty
    
    logger.info(f"MB52 stock dictionary created: {len(stock_dict)} unique materials")
    return stock_dict

def get_rm_requirements_for_fg(fg, fg_rm_df):
    """Get RM requirements for a given FG"""
    rm_req = {}
    fg_norm = str(fg).strip().upper()
    
    rows = fg_rm_df[fg_rm_df['Main_FG'].astype(str).str.strip().str.upper() == fg_norm]
    
    for _, row in rows.iterrows():
        rm_code = str(row['Raw Material']).strip()
        rm_desc = str(row['Raw Material Description']).strip()
        rm_qty = float(row['Raw Material qty']) if pd.notna(row['Raw Material qty']) else 0.0
        rm_uom = str(row['Raw Material UoM']).strip()
        
        if rm_code and rm_qty > 0:
            rm_req[rm_code] = (rm_desc, rm_qty, rm_uom)
    
    return rm_req

def process_fifo_allocation(df_orders, fg_rm_df, mb52_df):
    """FIFO ALLOCATION: Date-based FIFO across ALL orders"""
    logger.info("=== STARTING SHAREPOINT-BASED FIFO ALLOCATION PROCESSING ===")
    
    # Create initial MB52 stock dictionary for Plant 1000
    mb52_stock_dict = create_mb52_stock_dict(mb52_df)
    initial_stock_total = sum(mb52_stock_dict.values())
    
    # Initialize tracking variables
    results = []
    allocation_summary = {
        'total_orders': len(df_orders),
        'successful_orders': 0,
        'failed_orders': 0,
        'by_material': {},
        'stock_consumed': 0
    }
    
    logger.info(f"Processing {len(df_orders)} orders in strict DI date order (FIFO)")
    logger.info(f"Initial total stock available: {initial_stock_total:.3f} units across all materials")
    
    # Track statistics by material
    material_stats = {}
    
    # Process each order in DI date sequence
    for order_sequence, (original_idx, order) in enumerate(df_orders.iterrows(), 1):
        indus_part = str(order['Indus Part #']).strip().upper()
        order_qty = float(order['Qty'])
        order_di = order['DI']
        
        logger.info(f"\n--- Processing Order {order_sequence}/{len(df_orders)}: {indus_part} ---")
        logger.info(f"Order Details: Qty={order_qty}, DI={order_di}, Original Index={original_idx}")
        
        if indus_part not in material_stats:
            material_stats[indus_part] = {'successful': 0, 'total': 0}
        material_stats[indus_part]['total'] += 1
        
        # Get RM requirements for this FG
        rm_requirements = get_rm_requirements_for_fg(indus_part, fg_rm_df)
        
        if not rm_requirements:
            logger.warning(f"No RM requirements found for {indus_part} - marking as Error")
            results.append({
                'index': original_idx,
                'Indus Part #': indus_part,
                'Qty': order_qty,
                'DI': order_di,
                'Status': 'Error',
                'Shortage_Details': f'No RM requirements found in FG-RM Required sheet for {indus_part}'
            })
            continue
        
        logger.info(f"Found {len(rm_requirements)} RM components for {indus_part}")
        
        # Apply UOM filtering
        rm_requirements_filtered = {}
        excluded_count = 0
        
        for rm_code, (rm_desc, qty_per_fg, rm_uom) in rm_requirements.items():
            unit = str(rm_uom).strip().upper()
            if unit in ("L", "LT", "LITRE", "ROL", "ROLL"):
                excluded_count += 1
                logger.debug(f"Excluding {rm_code} due to UoM: {unit}")
                continue
            rm_requirements_filtered[rm_code] = (rm_desc, qty_per_fg, rm_uom)
        
        if excluded_count > 0:
            logger.info(f"Excluded {excluded_count} components due to UoM filtering")
        
        if not rm_requirements_filtered:
            logger.warning(f"No eligible RM components for {indus_part} after UoM filtering - marking as Error")
            results.append({
                'index': original_idx,
                'Indus Part #': indus_part,
                'Qty': order_qty,
                'DI': order_di,
                'Status': 'Error',
                'Shortage_Details': f'No eligible RM components (excluded UoMs: L/ROL) for {indus_part}'
            })
            continue
        
        logger.info(f"Processing {len(rm_requirements_filtered)} eligible RM components")
        
        # Calculate total RM requirements for this order
        can_fulfill_completely = True
        shortage_components = []
        required_materials = {}
        total_stock_needed = 0
        
        for rm_code, (rm_desc, qty_per_fg, rm_uom) in rm_requirements_filtered.items():
            total_required = qty_per_fg * order_qty
            available_stock = mb52_stock_dict.get(rm_code, 0)
            total_stock_needed += total_required
            
            required_materials[rm_code] = {
                'required': total_required,
                'available': available_stock,
                'desc': rm_desc,
                'uom': rm_uom
            }
            
            logger.debug(f" {rm_code}: Required={total_required:.3f} {rm_uom}, Available={available_stock:.3f}")
            
            if available_stock < total_required:
                can_fulfill_completely = False
                shortage_qty = total_required - available_stock
                shortage_components.append(f"{rm_desc} ({rm_code}): {format_shortage_quantity(shortage_qty, rm_uom)} {rm_uom}")
                logger.info(f" SHORTAGE: {rm_code} needs {total_required:.3f}, only {available_stock:.3f} available")
        
        logger.info(f"Total stock needed for order: {total_stock_needed:.3f} units")
        
        # CRITICAL FIFO LOGIC: Reserve stock regardless of shortage status
        stock_consumed_this_order = 0
        for rm_code, details in required_materials.items():
            stock_to_consume = min(details['required'], details['available'])
            if stock_to_consume > 0:
                mb52_stock_dict[rm_code] -= stock_to_consume
                stock_consumed_this_order += stock_to_consume
                logger.debug(f" Reserved {stock_to_consume:.3f} of {rm_code}, remaining: {mb52_stock_dict[rm_code]:.3f}")
        
        allocation_summary['stock_consumed'] += stock_consumed_this_order
        
        # Determine order status and record results
        if can_fulfill_completely:
            logger.info(f"Order FULFILLED: All RM components available, consumed {stock_consumed_this_order:.3f} units")
            results.append({
                'index': original_idx,
                'Indus Part #': indus_part,
                'Qty': order_qty,
                'DI': order_di,
                'Status': 'YES',
                'Shortage_Details': ''
            })
            allocation_summary['successful_orders'] += 1
            material_stats[indus_part]['successful'] += 1
        else:
            shortage_text = ", ".join(shortage_components)
            logger.info(f"Order FAILED: Component shortages detected, consumed {stock_consumed_this_order:.3f} units")
            logger.info(f"Shortage details: {shortage_text}")
            results.append({
                'index': original_idx,
                'Indus Part #': indus_part,
                'Qty': order_qty,
                'DI': order_di,
                'Status': 'NO',
                'Shortage_Details': shortage_text
            })
            allocation_summary['failed_orders'] += 1
        
        # Progress logging
        if order_sequence % 10 == 0 or order_sequence == len(df_orders):
            remaining_stock_total = sum(mb52_stock_dict.values())
            stock_utilization = ((initial_stock_total - remaining_stock_total) / initial_stock_total * 100) if initial_stock_total > 0 else 0
            logger.info(f"Progress: {order_sequence}/{len(df_orders)} orders processed, {stock_utilization:.1f}% stock utilized")
    
    # Compile final results
    allocation_summary['by_material'] = material_stats
    
    # Convert results to DataFrame and sort by original index
    results_df = pd.DataFrame(results)
    if not results_df.empty:
        results_df = results_df.sort_values('index').reset_index(drop=True)
    
    # Final summary logging
    final_stock_total = sum(mb52_stock_dict.values())
    total_consumed = initial_stock_total - final_stock_total
    stock_efficiency = (total_consumed / initial_stock_total * 100) if initial_stock_total > 0 else 0
    
    logger.info("\n=== SHAREPOINT-BASED FIFO ALLOCATION COMPLETED ===")
    logger.info(f"Stock Utilization Summary:")
    logger.info(f" Initial Stock: {initial_stock_total:.3f} units")
    logger.info(f" Consumed Stock: {total_consumed:.3f} units ({stock_efficiency:.1f}%)")
    logger.info(f" Remaining Stock: {final_stock_total:.3f} units")
    logger.info(f"Order Fulfillment Summary:")
    logger.info(f" Total Orders: {allocation_summary['total_orders']}")
    logger.info(f" Successful: {allocation_summary['successful_orders']} ({allocation_summary['successful_orders']/allocation_summary['total_orders']*100:.1f}%)")
    logger.info(f" Failed: {allocation_summary['failed_orders']} ({allocation_summary['failed_orders']/allocation_summary['total_orders']*100:.1f}%)")
    
    return results_df, allocation_summary

def format_shortage_quantity(qty, uom):
    """Return shortage quantity rounded per UOM: PC=1, KG=2, else=2"""
    try:
        qty = float(qty)
        if str(uom).strip().upper() == 'PC':
            return f"{qty:.1f}"
        elif str(uom).strip().upper() == 'KG':
            return f"{qty:.2f}"
        else:
            return f"{qty:.2f}"
    except Exception:
        return str(qty)

def validate_data_linkage(df_tracker, fg_rm_df):
    """Enhanced data validation with detailed logging"""
    logger.info("=== VALIDATING DATA LINKAGE ===")
    
    try:
        tracker_parts = set(df_tracker['Indus Part #'].astype(str).str.strip().str.upper().unique())
        fg_rm_parts = set(fg_rm_df['Main_FG'].astype(str).str.strip().str.upper().unique())
        
        logger.info(f"Production Tracker unique parts: {len(tracker_parts)}")
        logger.info(f"FG-RM Required unique parts: {len(fg_rm_parts)}")
        
        # Find matches and mismatches
        matches = tracker_parts.intersection(fg_rm_parts)
        tracker_only = tracker_parts - fg_rm_parts
        fg_rm_only = fg_rm_parts - tracker_parts
        
        logger.info(f"Matching parts: {len(matches)}")
        logger.info(f"Tracker-only parts: {len(tracker_only)}")
        logger.info(f"FG-RM-only parts: {len(fg_rm_only)}")
        
        if matches:
            logger.info(f"Sample matches: {list(matches)[:5]}")
        if tracker_only:
            logger.warning(f"Parts in Tracker but not in FG-RM Required: {list(tracker_only)[:5]}")
        
        return len(matches) > 0, matches
        
    except Exception as e:
        logger.error(f"Data validation error: {e}")
        return False, set()

# --- MAIN FUNCTION: SHAREPOINT-ONLY PROCESSING ---
def main():
    """COMPLETE SHAREPOINT-ONLY BOM layer-wise FIFO allocation system"""
    logger.info("Starting SHAREPOINT-ONLY BOM layer-wise FIFO allocation system...")
    start_time = time.time()
    
    # SharePoint configuration
    SITE_ID = "201c767f-23cc-4eff-b76c-47a275706531,094b21f9-6ed4-47c3-a751-3b217a18b98e"
    SHAREPOINT_FILE = "Manufacturing Plan New.xlsx"
    
    try:
        # === PHASE 1: DATA COLLECTION ===
        logger.info("=== DATA COLLECTION PHASE ===")
        
        # Enhanced concurrent data fetching
        df_bom = fetch_bom_data(SAP_API_URL, USERNAME, PASSWORD, VERIFY_SSL, PREFERRED_FORMAT, BATCH_SIZE, max_workers=15)
        df_bom = standardize_bom_data(df_bom)
        
        # Fetch other data concurrently
        with ThreadPoolExecutor(max_workers=2) as executor:
            header_future = executor.submit(fetch_material_bom_headers, SAP_API_URL, USERNAME, PASSWORD, VERIFY_SSL, PREFERRED_FORMAT, BATCH_SIZE)
            stock_future = executor.submit(fetch_material_stock_data, SAP_API_URL, USERNAME, PASSWORD, VERIFY_SSL, PREFERRED_FORMAT, BATCH_SIZE)
            
            header_df = header_future.result()
            stock_df = stock_future.result()
        
        header_mapping = dict(zip(header_df["Material"], header_df["ProductDescription"]))
        header_qty_mapping = dict(zip(header_df["Material"], pd.to_numeric(header_df["BOMHeaderQuantityInBaseUnit"], errors="coerce").fillna(1)))
        
        logger.info(f"Data collection completed - BOM: {len(df_bom)}, Headers: {len(header_df)}, Stock: {len(stock_df)}")
        
        # === PHASE 2: BOM TREE STRUCTURE CREATION ===
        logger.info("=== BOM TREE STRUCTURE CREATION ===")
        
        # Create BOM tree structure (in memory, no local files)
        logger.info("Creating BOM layer-wise structure...")
        all_paths = []
        for fg in df_bom["Material"].unique():
            fg_paths = list(tree_explode(df_bom, fg, MAX_LAYERS))
            all_paths.extend(fg_paths)
        
        logger.info(f"BOM tree structure created with {len(all_paths)} paths")
        
        # === PHASE 3: SHAREPOINT EXCEL SESSION CREATION ===
        logger.info("=== SHAREPOINT EXCEL SESSION CREATION ===")
        
        # Get access token and create session
        access_token = get_access_token()
        session_id = None
        
        # Create SharePoint Excel session
        for session_attempt in range(2):
            try:
                session_url = f"https://graph.microsoft.com/v1.0/sites/{SITE_ID}/drive/root:/{SHAREPOINT_FILE}:/workbook/createSession"
                session_headers = {
                    "Authorization": f"Bearer {access_token}",
                    "Content-Type": "application/json"
                }
                session_payload = {"persistChanges": True}
                
                session_response = requests.post(session_url, headers=session_headers, json=session_payload, timeout=30)
                if session_response.status_code in [200, 201]:
                    session_id = session_response.json().get("id")
                    logger.info(f"Created SharePoint Excel session: {session_id}")
                    break
                else:
                    if session_attempt < 1:
                        time.sleep(2)
                        continue
                    else:
                        logger.error("Failed to create SharePoint session")
                        return
            except Exception as session_error:
                if session_attempt < 1:
                    time.sleep(2)
                    continue
                else:
                    logger.error(f"Failed to create SharePoint session: {session_error}")
                    return
        
        if not session_id:
            logger.error("Could not create SharePoint Excel session")
            return
            
        try:
            # === PHASE 4: SHAREPOINT SHEET CREATION ===
            logger.info("=== SHAREPOINT SHEET CREATION ===")
            
            # Create BOM Layer-wise sheet directly in SharePoint
            logger.info("Creating BOM Layer-wise sheet in SharePoint...")
            success_main = create_sharepoint_bom_layerwise_sheet(SITE_ID, SHAREPOINT_FILE, session_id, access_token, df_bom, all_paths, MAX_LAYERS, header_mapping, header_qty_mapping)
            time.sleep(2)
            
            # Read the created BOM Layer-wise data for FG-RM processing
            logger.info("Reading created BOM Layer-wise data...")
            df_main_sheet = read_sharepoint_worksheet_data(SITE_ID, SHAREPOINT_FILE, "BOM Layer-wise", session_id, access_token)
            
            # Create FG-RM Required sheet
            logger.info("Creating FG-RM Required sheet in SharePoint...")
            success_fg_rm = create_sharepoint_fg_rm_sheet(SITE_ID, SHAREPOINT_FILE, session_id, access_token, df_main_sheet, MAX_LAYERS)
            time.sleep(2)
            
            # Create MB52 Raw Data sheet
            logger.info("Creating MB52 Raw Data sheet in SharePoint...")
            success_stock = create_sharepoint_mb52_sheet(SITE_ID, SHAREPOINT_FILE, session_id, access_token, stock_df)
            time.sleep(1)
            
            # === PHASE 5: SHAREPOINT DATA VALIDATION ===
            logger.info("=== SHAREPOINT DATA VALIDATION ===")
            
            # Read Production Tracker data from SharePoint
            df_orders, full_df = read_sharepoint_production_tracker(SITE_ID, SHAREPOINT_FILE, session_id, access_token)
            
            if df_orders.empty:
                logger.error("No SharePoint Production Tracker data available for processing")
                return
            
            logger.info(f"Ready to process {len(df_orders)} orders from SharePoint Production Tracker")
            
            # Read FG-RM Required and MB52 data from SharePoint
            fg_rm_df = read_sharepoint_worksheet_data(SITE_ID, SHAREPOINT_FILE, "FG-RM Required", session_id, access_token)
            mb52_df = read_sharepoint_worksheet_data(SITE_ID, SHAREPOINT_FILE, "MB52 Raw Data", session_id, access_token)
            
            # Validate data linkage
            is_valid, matches = validate_data_linkage(df_orders, fg_rm_df)
            if not is_valid:
                logger.error("CRITICAL: No data linkage between Production Tracker and FG-RM Required!")
                return
            
            logger.info(f"Data validation passed: {len(matches)} matching materials")
            
            # === PHASE 6: SHAREPOINT FIFO ALLOCATION EXECUTION ===
            logger.info("=== SHAREPOINT FIFO ALLOCATION EXECUTION ===")
            
            # Execute the core FIFO allocation logic using SharePoint data
            results_df, allocation_summary = process_fifo_allocation(df_orders, fg_rm_df, mb52_df)
            
            # === PHASE 7: SHAREPOINT RESULTS UPDATE ===
            logger.info("=== SHAREPOINT RESULTS UPDATE ===")
            
            # Update SharePoint Production Tracker with results
            success_tracker = update_sharepoint_production_tracker_results(SITE_ID, SHAREPOINT_FILE, session_id, access_token, results_df, full_df)
            
            # Set timestamps in all sheets
            set_sharepoint_timestamp(SITE_ID, SHAREPOINT_FILE, "Production Tracker", session_id, access_token, "R1")
            set_sharepoint_timestamp(SITE_ID, SHAREPOINT_FILE, "BOM Layer-wise", session_id, access_token, "S1")
            set_sharepoint_timestamp(SITE_ID, SHAREPOINT_FILE, "FG-RM Required", session_id, access_token, "F1")
            
            # Log SharePoint operations results
            if success_main and success_fg_rm and success_stock and success_tracker:
                logger.info("All SharePoint Excel operations completed successfully")
            else:
                logger.warning("Some SharePoint Excel operations failed")
                
        finally:
            # Always close the session
            try:
                close_session_url = f"https://graph.microsoft.com/v1.0/sites/{SITE_ID}/drive/root:/{SHAREPOINT_FILE}:/workbook/closeSession"
                close_headers = {
                    "Authorization": f"Bearer {access_token}",
                    "workbook-session-id": session_id
                }
                close_response = requests.post(close_session_url, headers=close_headers, timeout=15)
                if close_response.status_code in [200, 204]:
                    logger.info("SharePoint Excel session closed successfully")
            except Exception as close_error:
                logger.warning(f"Error closing session: {close_error}")
        
        # === PHASE 8: FINAL SUMMARY REPORT ===
        logger.info("=== SHAREPOINT-ONLY ALLOCATION SUMMARY REPORT ===")
        logger.info(f" Total Orders: {allocation_summary['total_orders']}")
        logger.info(f" Successful: {allocation_summary['successful_orders']} ({allocation_summary['successful_orders']/allocation_summary['total_orders']*100:.1f}%)")
        logger.info(f" Failed: {allocation_summary['failed_orders']} ({allocation_summary['failed_orders']/allocation_summary['total_orders']*100:.1f}%)")
        
        if not results_df.empty:
            total_qty_requested = results_df['Qty'].sum()
            successful_qty = results_df[results_df['Status'] == 'YES']['Qty'].sum()
            failed_qty = results_df[results_df['Status'] == 'NO']['Qty'].sum()
            logger.info(f" Total Quantity Requested: {total_qty_requested}")
            logger.info(f" Quantity Allocated: {successful_qty}")
            logger.info(f" Quantity Short: {failed_qty}")
        
        elapsed_time = time.time() - start_time
        logger.info(f"SHAREPOINT-ONLY FIFO allocation system completed successfully in {elapsed_time:.2f} seconds")
        
    except Exception as e:
        logger.error(f"Critical error in main process: {e}")
        import traceback
        logger.error(f"Full traceback: {traceback.format_exc()}")
        raise

if __name__ == "__main__":
    main()


