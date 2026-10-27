import requests
import pandas as pd
import xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Border, Side, Alignment, PatternFill, Font
import logging
import warnings
import time
import os
import sys
from openpyxl import load_workbook
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
import msal

if sys.platform == "win32":
    import os
    os.environ["PYTHONIOENCODING"] = "utf-8"

warnings.filterwarnings('ignore')

# --- Logging Setup ---
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s %(levelname)s %(message)s',
    handlers=[
        logging.FileHandler('bom_layerwise.log'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# --- Configuration ---
SAP_API_URL = "https://my409486-api.s4hana.cloud.sap/sap/opu/odata/sap/API_BILL_OF_MATERIAL_SRV;v=0002"
USERNAME = (os.environ.get("SAP_USERNAME") or "").strip()
PASSWORD = (os.environ.get("SAP_PASSWORD") or "").strip()
if not USERNAME or not PASSWORD:
    raise RuntimeError(
        "SAP_USERNAME / SAP_PASSWORD not set - add them as GitHub repository secrets"
    )
VERIFY_SSL = True
PREFERRED_FORMAT = 'json'
BATCH_SIZE = 500
MAX_LAYERS = 4

import re

# --- Material group code -> description (static list supplied by the plant) ---
_MATERIAL_GROUP_TEXT = """
1101|Brass parts/fittings
1102|Copper tube
1103|Aluminium tube
1104|Aluminium foil
1105|Sheet Metal
1106|Rubber Plug
1107|Distributor
1108|Components
1109|Tripod, U Bend
1110|Brazing rings
1111|Header Aluminium
1201|Tooling spares
1202|Mech-Spares
1203|Elect-Spares
1204|Hydraulic Spares
1205|Pnuematic Spares
1206|Maintenance consumab
1207|Jigs & fixtures
1301|Chemicals, Paints
1302|Lubricants all types
1303|Industrial Gases
1304|Brazing rings, rods
1305|Stationary
1306|IT Consumables
1307|Maintenance consumab
1308|Factory consumables
1401|Plant Machinery, Equ
1402|Furniture & Fixtures
1403|Vehicles
1404|Office Equipement
1405|Building
1406|Software
1501|Admin & IT Service
1502|Material Handling Se
1503|Waste disposal servi
1504|Technical Services
1505|Professional fees
1506|Freight & forwarding
1601|Aluminum Scrap
1602|Copper Scrap
1603|GI and MS Scrap
1604|Brass scrap
1605|X
1606|Empty Drums
1607|SS Scrap
1701|Endplate
1702|Copper hair pins
1703|Aluminium hair pins
1704|Expanded Block
1705|Header Copper
1706|Header Aluminium
1707|Shroud and Bracket
1708|Return Bend
1709|Finished Component
1710|Subcontracted
1711|Combi EB
1712|Combi FC
1801|Outer Packing Materi
1802|Wooden Material
1803|Packing consumables
1804|Filler material
1901|Evaporator
1902|Condenser
1903|Heater
1904|Combi
1905|Assembly
1906|Sales Kit
2001|Product Inquiry
2101|Header Assy Copper
2102|Header Assy Aluminiu
2201|Coil scrap internal
2202|Coil scrap external
A001|On Account Billed
A002|On Account Utilized
B001|One Time Credit
B002|Cloud Revenue
E001|Accommodation
E002|Airfare
E003|Ground Transp.
E004|Travel Expenses misc
L001|Trading Materials
L002|Raw Materials
L003|Semi-Finished Goods
L004|Finished Goods
L005|Unvaluated Materials
L006|Returnable Packaging
L007|Licenses
L008|Hardware
P000|Contract Type
P001|Services
P002|Expenses
P003|Periodic Service
YBFA01|Real Estate (Land)
YBFA02|Buildings
YBFA03|Land Improvements
YBFA04|Leasehold Improvemen
YBFA05|Machinery Equipment
YBFA06|Fixtures Fittings
YBFA07|Vehicles
YBFA08|Computer Hardware
YBFA09|Computer Software
YBFA10|Low Value Assets
YBFA11|Other Intangibles
YBFA12|Office Equipment
YBMM00|Non-Sto Mat. w.o ID
YBMM01|Non-Sto Mat. w. ID
YBPM01|Spare parts
YBPM02|Spare parts services
"""
MATERIAL_GROUP_DESC = {}
for _line in _MATERIAL_GROUP_TEXT.strip().splitlines():
    _code, _desc = _line.split("|", 1)
    MATERIAL_GROUP_DESC[_code.strip().upper()] = _desc.strip()

def material_group_desc(code):
    """Description for a material group code; blank if the code is blank or not in the list."""
    return MATERIAL_GROUP_DESC.get(str(code).strip().upper(), "")

def is_foil_component(desc, material_group):
    """True for an aluminium foil line (material group 1104 or an 'AL FOIL' description), not the cut-piece scrap line."""
    d = str(desc).upper().strip()
    if re.search(r"CUT ?PIECE", d):
        return False
    return str(material_group).strip() == "1104" or bool(re.match(r"^(AL|ALUMINI?UM)\s+FOIL", d))

def _fmt_num(v):
    if v is None:
        return ""
    return str(int(v)) if float(v).is_integer() else str(round(float(v), 3))

# --- Fin-pack text parser (handles the messy typed text on the foil item text) ---

def _num(m):
    if not m:
        return None
    v = next((g for g in m.groups() if g), None)
    try:
        return float(v) if v is not None else None
    except ValueError:
        return None

_RX_FPI = re.compile(r"FPI[-\s]*0?(\d+(?:\.\d+)?)|(\d+(?:\.\d+)?)\.?FPI", re.I)
_RX_FIN = re.compile(r"(?<![A-Z])FINS?[-\s]*(\d+(?:\.\d+)?)|(\d+(?:\.\d+)?)FINS?", re.I)
_RX_FL = re.compile(r"F/?L[-\s]*(\d+(?:\.\d+)?)", re.I)

_EDGE_RIPPLE = re.compile(r"RI+P\w*")
_EDGE_STRAIGHT = re.compile(r"STR[AI]*G?[IH]*T\w*|STRAIGH\w*|STRAGIHT|STRIGHT")


def _classify_surface(s):
    s = re.sub(r"[^A-Z0-9 ]", " ", s.upper())
    s = re.sub(r"\s+", " ", s).strip()
    if not s:
        return ""
    if "TSUNAMI" in s:
        return "Tsunami"
    lanced = re.search(r"L[AE]NC", s)
    raised = re.search(r"RAI", s) or re.search(r"(^|\s)R\s?(?=L[AE]NC)", s) or (re.search(r"RAIS", s))
    louvered = re.search(r"L[O0]U?V|LOUER|LOVER", s)
    flat = re.search(r"\bFLAT\b|\bFLA\b|FLAT", s)
    if raised and (lanced or re.search(r"\bL\b|\bLA\b", s)):
        return "Raised lanced"
    if lanced:
        return "Flat lanced" if flat else "Lanced"
    if louvered:
        return "Flat louvered" if flat else "Louvered"
    if re.search(r"SIN[EW]|\bSW\b|WAVY", s):
        return "Sinewave"
    if re.search(r"COR+U|\bCR\b", s):
        return "Corrugated"
    if flat:
        return "Flat"
    return ""


def parse_fin_spec(text):
    """Return dict: fpi, fins, length_mm, surface, edge, check. Never raises."""
    out = {"fpi": None, "fins": None, "length_mm": None, "surface": "", "edge": "", "check": ""}
    raw = "" if text is None else str(text).strip()
    if not raw:
        out["check"] = "No fin text"
        return out
    try:
        s = raw.upper()
        s = re.sub(r"\.{2,}", ".", s)                       # 122..16 -> 122.16
        s = re.sub(r"\([^)]*\)", " ", s)                     # (-0.5)
        out["fpi"] = _num(_RX_FPI.search(s))
        out["fins"] = _num(_RX_FIN.search(s))
        out["length_mm"] = _num(_RX_FL.search(s))

        # remove the numeric spec so only the pattern words are left
        t = _RX_FPI.sub(" ", s)
        t = _RX_FIN.sub(" ", t)
        t = re.sub(r"F/?L[-\s]*\d+(?:\.\d+)?\s*(?:MM)?", " ", t)
        t = re.sub(r",\s*[\d\.]+\s*(?:\"|''|\"\"|IN\b|INCH\b)", " ", t)   # , 25"  or , 7''
        t = re.sub(r"[\d\.]+\s*(?:\"|'')", " ", t)
        t = re.sub(r"HYDROPH\w*|HYDR\b|HYDRO\b|\bBARE\b", " ", t)
        t = re.sub(r"\bMM\b", " ", t)
        t = re.sub(r"[,;]", " ", t)
        t = re.sub(r"\s+", " ", t).strip(" -")

        parts = [p.strip() for p in re.split(r"\s*/\s*", t) if p.strip()]
        surface_parts, edge = [], ""
        for idx, p in enumerate(parts):
            m_r, m_s = _EDGE_RIPPLE.search(p), _EDGE_STRAIGHT.search(p)
            if m_r or m_s:
                m = m_r or m_s
                edge = edge or ("Ripple" if m_r else "Straight")
                rest = (p[:m.start()] + " " + p[m.end():]).strip()
                if rest:
                    surface_parts.append(rest)
            elif re.fullmatch(r"R|RI", p):
                edge = edge or "Ripple"
            elif "PLAIN" in p:
                edge = edge or "Plain"
            elif idx > 0 and surface_parts and re.fullmatch(r"FLAT(\s+EDGE)?|FLA", p):
                edge = edge or "Flat"
            elif idx > 0 and surface_parts and re.search(r"SIN[EW]|\bSW\b", p) and \
                    _classify_surface(" ".join(surface_parts)) not in ("", "Sinewave"):
                edge = edge or "Sinewave"
            else:
                surface_parts.append(p)
        surface = _classify_surface(" ".join(surface_parts))
        if not surface and re.fullmatch(r"S", " ".join(surface_parts).strip()) and edge == "Ripple":
            surface = "Sinewave"
        out["surface"], out["edge"] = surface, edge

        if out["fpi"] is None or out["fins"] is None or out["length_mm"] is None:
            out["check"] = "Could not read FPI / fins / length"
        else:
            calc = out["fpi"] * out["length_mm"] / 25.4
            if abs(calc - out["fins"]) > max(2.0, 0.03 * out["fins"]):
                out["check"] = "Fins do not match FPI x length"
        if not out["check"] and not surface:
            out["check"] = "Surface pattern not stated"
    except Exception:
        out["check"] = "Could not parse"
    return out


# === SHAREPOINT-ONLY GRAPH API INTEGRATION ===
import msal
import json

class AuthConfig:
    CLIENT_ID = (os.environ.get("AZURE_CLIENT_ID") or "").strip()
    CLIENT_SECRET = (os.environ.get("AZURE_CLIENT_SECRET") or "").strip()
    AUTHORITY = 'https://login.microsoftonline.com/ed97e9bb-e119-4bd4-ab00-307d64bdf908'
    if not CLIENT_ID or not CLIENT_SECRET:
        raise RuntimeError(
            "AZURE_CLIENT_ID / AZURE_CLIENT_SECRET not set - add them as GitHub repository secrets"
        )

def get_access_token():
    """Get Microsoft Graph API access token"""
    auth_config = AuthConfig()
    app = msal.ConfidentialClientApplication(
        auth_config.CLIENT_ID, 
        authority=auth_config.AUTHORITY, 
        client_credential=auth_config.CLIENT_SECRET
    )
    
    result = app.acquire_token_silent(scopes=['https://graph.microsoft.com/.default'], account=None)
    if not result:
        result = app.acquire_token_for_client(scopes=['https://graph.microsoft.com/.default'])
    
    if 'access_token' in result:
        return result['access_token']
    else:
        raise Exception(f"Failed to acquire access token: {result.get('error_description', result)}")

def get_excel_column_letter(col_num):
    """Convert column number to Excel column letter (e.g., 1 -> A, 27 -> AA)"""
    letter = ""
    while col_num > 0:
        col_num -= 1
        letter = chr(col_num % 26 + ord('A')) + letter
        col_num //= 26
    return letter

# --- Excel session handling: if SharePoint drops the session (for example because someone else edited the
# --- workbook at the same moment), wait, open a new session and write the whole sheet again.
ACTIVE_SESSION = {"id": None}

def _create_workbook_session(site_id, file_server_relative_url, access_token):
    url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/createSession"
    try:
        r = requests.post(url, headers={"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"},
                          json={"persistChanges": True}, timeout=30)
        if r.status_code in (200, 201):
            return r.json().get("id")
        logger.error(f"Could not open a new Excel session: {r.status_code} {r.text[:200]}")
    except Exception as e:
        logger.error(f"Could not open a new Excel session: {e}")
    return None

def create_sharepoint_worksheet_optimized(site_id, file_server_relative_url, worksheet_name, df, access_token, session_id, *args, **kwargs):
    """Write a worksheet. If the Excel session is dropped mid-write, open a fresh session and rewrite the whole sheet."""
    max_sheet_attempts = 4
    current = ACTIVE_SESSION["id"] or session_id
    for attempt in range(1, max_sheet_attempts + 1):
        result = _write_worksheet_once(site_id, file_server_relative_url, worksheet_name, df, access_token, current, *args, **kwargs)
        if result != "INVALID_SESSION":
            return result
        logger.warning(f"Excel session was dropped while writing '{worksheet_name}' (attempt {attempt}/{max_sheet_attempts}). "
                       f"This usually means someone else has the workbook open and is editing it.")
        if attempt == max_sheet_attempts:
            break
        time.sleep(30 * attempt)
        new_id = _create_workbook_session(site_id, file_server_relative_url, access_token)
        if new_id:
            current = new_id
            ACTIVE_SESSION["id"] = new_id
            logger.info(f"Opened a new Excel session for '{worksheet_name}' and starting that sheet again")
    logger.error(f"Giving up on '{worksheet_name}' after {max_sheet_attempts} attempts - close the workbook everywhere and run again")
    return False

def _write_worksheet_once(site_id, file_server_relative_url, worksheet_name, df, access_token, session_id, chunk_size=1000, max_retries=5, native_types=False):
    """Robust creation/update for large SharePoint sheets with data validation and exponential backoff"""
    session = requests.Session()
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": ACTIVE_SESSION["id"] or session_id
    }

    logger.info(f"Creating/updating SharePoint worksheet: {worksheet_name}")
    # Check for worksheet existence, create if needed (as before)
    ws_list_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets"
    ws_list_resp = session.get(ws_list_url, headers=api_headers)
    if ws_list_resp.status_code != 200 and "InvalidSession" in ws_list_resp.text:
        return "INVALID_SESSION"
    ws_names = [ws['name'] for ws in ws_list_resp.json().get('value', [])] if ws_list_resp.status_code == 200 else []
    if worksheet_name not in ws_names:
        create_ws_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets/add"
        session.post(create_ws_url, headers=api_headers, json={"name": worksheet_name})
        time.sleep(2)
    # Clear data in full
    try:
        clear_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/usedRange/clear"
        session.post(clear_url, headers=api_headers, json={"applyTo": "Contents"}, timeout=60)
        time.sleep(1)
    except Exception as clear_error:
        logger.warning(f"Clear operation failed: {clear_error}")

    # Write headers first
    headers = df.columns.tolist()
    if native_types:
        data_conv = df.astype(object).where(~df.isna(), "")  # keep numbers as numbers
    else:
        data_conv = df.astype(str).where(~df.isna(), "")  # Ensures all strings
    num_cols = len(headers)
    header_range = f"A1:{get_excel_column_letter(num_cols)}1"
    header_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{header_range}')"
    session.patch(header_url, headers=api_headers, json={"values": [headers]}, timeout=60)

    # Chunked rows
    total_rows = len(df)
    total_chunks = (total_rows + chunk_size - 1) // chunk_size
    for chunk_idx in range(0, total_rows, chunk_size):
        chunk_end = min(chunk_idx + chunk_size, total_rows)
        chunk_data = data_conv.iloc[chunk_idx:chunk_end].values.tolist()
        chunk_range = f"A{chunk_idx+2}:{get_excel_column_letter(num_cols)}{chunk_end+1}"
        chunk_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{chunk_range}')"
        # Exponential backoff
        for attempt in range(max_retries):
            try:
                chunk_response = session.patch(
                    chunk_url, headers=api_headers, json={"values": chunk_data}, timeout=120
                )
                if chunk_response.status_code in [200, 201]:
                    logger.info(f"Chunk {(chunk_idx//chunk_size)+1}/{total_chunks} successful for '{worksheet_name}' (rows {chunk_idx+2}-{chunk_end+1})")
                    break
                elif chunk_response.status_code in {429, 503, 504}:
                    wait = 2 ** attempt
                    logger.warning(f"Throttling or timeout on chunk {(chunk_idx//chunk_size)+1}, wait {wait}s")
                    time.sleep(wait)
                elif "InvalidSession" in chunk_response.text:
                    logger.warning(f"Excel session dropped on chunk {(chunk_idx//chunk_size)+1} of '{worksheet_name}'")
                    return "INVALID_SESSION"
                else:
                    if attempt == max_retries - 1:
                        logger.error(f"Chunk {(chunk_idx//chunk_size)+1} for '{worksheet_name}' failed after {max_retries} attempts: {chunk_response.text}")
                    else:
                        time.sleep(2)
            except Exception as ex:
                wait = 2 ** attempt
                logger.warning(f"Exception on chunk {(chunk_idx//chunk_size)+1}, attempt {attempt+1}: {ex}, waiting {wait}s")
                time.sleep(wait)
        else:
            logger.error(f"Giving up on chunk {(chunk_idx//chunk_size)+1} for '{worksheet_name}'")
            return False
    logger.info(f"'{worksheet_name}' data upload complete.")
    apply_sharepoint_formatting_fixed(site_id, file_server_relative_url, worksheet_name, session_id, access_token)
    return True
def apply_sharepoint_formatting_fixed(site_id, file_server_relative_url, worksheet_name, session_id, access_token):
    """Apply formatting to SharePoint worksheet with proper border coverage"""
    session = requests.Session()
    
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": ACTIVE_SESSION["id"] or session_id
    }
    
    try:
        # Get the actual used range from the server
        used_range_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/usedRange"
        resp = session.get(used_range_url, headers=api_headers)
        if resp.status_code != 200:
            logger.warning(f"Could not retrieve usedRange for '{worksheet_name}': {resp.status_code}")
            return
        
        address = resp.json().get("address", "")
        if "!" in address:
            address = address.split("!", 1)[1]
        
        # Apply right alignment for all data
        align_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{address}')/format"
        align_payload = {"horizontalAlignment": "Right"}
        session.patch(align_url, headers=api_headers, json=align_payload)
        
        # Apply borders to entire used range
        border_ids = ["EdgeTop", "EdgeBottom", "EdgeLeft", "EdgeRight", "InsideHorizontal", "InsideVertical"]
        for border_id in border_ids:
            border_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{address}')/format/borders('{border_id}')"
            border_payload = {"style": "Continuous", "color": "#000000", "weight": "Thin"}
            session.patch(border_url, headers=api_headers, json=border_payload)
        
        # Get header range
        if address:
            cols = address.split(':')
            if len(cols) == 2:
                start_col = re.match(r"[A-Z]+", cols[0]).group(0)  # first column letter(s)
                end_col = re.match(r"[A-Z]+", cols[1]).group(0)    # last column letter(s)
                header_range = f"{start_col}1:{end_col}1"
                
                # Format header row
                header_fill_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{header_range}')/format/fill"
                header_fill_payload = {"color": "#366092"}
                session.patch(header_fill_url, headers=api_headers, json=header_fill_payload)
                
                header_font_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{header_range}')/format/font"
                header_font_payload = {"bold": True, "color": "#FFFFFF"}
                session.patch(header_font_url, headers=api_headers, json=header_font_payload)
        
        # Auto-fit columns
        autofit_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{address}')/format/autofitColumns"
        session.post(autofit_url, headers=api_headers)
        
        logger.info(f"Applied formatting with full border coverage to '{worksheet_name}' (range: {address})")
        
    except Exception as e:
        logger.warning(f"Error applying formatting to '{worksheet_name}': {e}")

def create_sharepoint_bom_layerwise_sheet(site_id, file_server_relative_url, session_id, access_token, df_bom, all_paths, max_layers, header_mapping, header_qty_mapping):
    """Create BOM Layer-wise sheet directly in SharePoint"""
    logger.info("Creating BOM Layer-wise sheet directly in SharePoint...")

    # Same processing logic as local export_tree_paths
    columns = ["Main_FG", "Main_FG_Description"]
    for lvl in range(1, max_layers+1):
        columns.extend([
            f"BOM_Layer_{lvl}",
            f"BOM_Layer_{lvl}_Desc",
            f"BOM_Layer_{lvl}_Quantity",
            f"BOM_Layer_{lvl}_Unit"
        ])

    # Item Text (line 1) columns go at the far right, after all existing columns
    for lvl in range(1, max_layers+1):
        columns.append(f"BOM_Layer_{lvl}_ItemText")
    for lvl in range(1, max_layers+1):
        columns.extend([f"BOM_Layer_{lvl}_MaterialGroup", f"BOM_Layer_{lvl}_MaterialGroupDesc"])
    columns.extend(["Fin_FPI", "Fin_Fins_Per_Pack", "Fin_Length_mm", "Fin_Surface", "Fin_Edge", "Fin_Spec_Check"])

    records = []
    for path in all_paths:
        row = dict()
        if path:
            first = path[0]
            row["Main_FG"] = first[0]
            row["Main_FG_Description"] = header_mapping.get(first[0], "")
        else:
            row["Main_FG"], row["Main_FG_Description"] = "", ""

        for i, comp in enumerate(path):
            if i < max_layers:
                row[f"BOM_Layer_{i+1}"] = comp[4]
                row[f"BOM_Layer_{i+1}_Desc"] = comp[1]
                row[f"BOM_Layer_{i+1}_Quantity"] = comp[2]
                row[f"BOM_Layer_{i+1}_Unit"] = comp[3]
                row[f"BOM_Layer_{i+1}_ItemText"] = comp[6]
                row[f"BOM_Layer_{i+1}_MaterialGroup"] = comp[7]
                row[f"BOM_Layer_{i+1}_MaterialGroupDesc"] = material_group_desc(comp[7])

        # fin-pack spec: read from the item text of the aluminium foil line in this path
        for i, comp in enumerate(path):
            if i < max_layers and is_foil_component(comp[1], comp[7]):
                fin = parse_fin_spec(comp[6])
                row["Fin_FPI"] = _fmt_num(fin["fpi"])
                row["Fin_Fins_Per_Pack"] = _fmt_num(fin["fins"])
                row["Fin_Length_mm"] = _fmt_num(fin["length_mm"])
                row["Fin_Surface"] = fin["surface"]
                row["Fin_Edge"] = fin["edge"]
                row["Fin_Spec_Check"] = fin["check"]
                break
        records.append(row)

    df_out = pd.DataFrame(records, columns=columns)

    # Apply same filtering as local
    filtered_df_out = df_out[df_out["Main_FG"].str.startswith(("A", "C", "I"))].copy()
    logger.info(f"Filtered rows (A/C/I only), remaining count: {len(filtered_df_out)}")

    # Apply BOM header quantity division
    def divide_qtys_by_divisor(qty_str, divisor):
        try:
            if not isinstance(qty_str, str):
                qty_str = str(qty_str) if qty_str is not None else ''
            qtys = [float(q) if q.strip() else 0 for q in qty_str.split(" | ")]
            divided = [q / divisor if divisor != 0 else q for q in qtys]
            result = " | ".join([str(round(q, 10)).rstrip('0').rstrip('.') if q != 0 else '0' for q in divided])
            return result
        except Exception as e:
            return qty_str

    for main_fg, group in filtered_df_out.groupby("Main_FG"):
        divisor = header_qty_mapping.get(main_fg, 1)
        try:
            divisor = float(divisor)
        except Exception:
            divisor = 1

        if pd.notna(divisor) and divisor != 0:
            mask = filtered_df_out["Main_FG"] == main_fg
            # ensure column is object/string before assigning string results
            filtered_df_out["BOM_Layer_1_Quantity"] = filtered_df_out["BOM_Layer_1_Quantity"].astype("object")
            filtered_df_out.loc[mask, "BOM_Layer_1_Quantity"] = (
                filtered_df_out.loc[mask, "BOM_Layer_1_Quantity"]
                .apply(divide_qtys_by_divisor, divisor=divisor)
            )

    # Convert G to KG for all layers
    for lvl in range(1, max_layers + 1):
        qty_col = f"BOM_Layer_{lvl}_Quantity"
        unit_col = f"BOM_Layer_{lvl}_Unit"

        if qty_col in filtered_df_out.columns and unit_col in filtered_df_out.columns:
            mask = filtered_df_out[unit_col].astype(str).str.strip().str.upper() == 'G'
            n_g = mask.sum()

            if n_g > 0:
                def convert_and_scale_g_to_kg(row):
                    qty_value = row[qty_col]
                    try:
                        q = float(qty_value)
                    except Exception:
                        q = 0.0
                    return q / 1000

                filtered_df_out.loc[mask, qty_col] = filtered_df_out.loc[mask].apply(convert_and_scale_g_to_kg, axis=1)
                filtered_df_out.loc[mask, unit_col] = 'KG'

    return create_sharepoint_worksheet_optimized(site_id, file_server_relative_url, "BOM Layer-wise", filtered_df_out, access_token, session_id)

def create_sharepoint_fg_rm_sheet(site_id, file_server_relative_url, session_id, access_token, main_sheet_df, max_layers):
	"""Create FG-RM Required sheet directly in SharePoint with correct cascading (no scaling for Layer-1)."""

	logger.info("Creating FG-RM Required sheet directly in SharePoint...")

	summary_records = []

	for main_fg in main_sheet_df["Main_FG"].unique():
		if not isinstance(main_fg, str) or not main_fg:
			continue

		fg_rows = main_sheet_df[main_sheet_df["Main_FG"] == main_fg]
		rm_dict = {}

		for _, row in fg_rows.iterrows():
			# Base Layer-1 quantity (used only as multiplier for deeper layers)
			layer_1_qty_str = str(row.get("BOM_Layer_1_Quantity", "1")).strip()
			try:
				layer_1_qtys = [float(q) if q.strip() else 1 for q in layer_1_qty_str.split(" | ")]
				layer_1_total_qty = sum(layer_1_qtys) if layer_1_qtys else 1
			except Exception:
				layer_1_total_qty = 1

			for lvl in range(1, max_layers + 1):
				mat_col = f"BOM_Layer_{lvl}"
				desc_col = f"BOM_Layer_{lvl}_Desc"
				qty_col = f"BOM_Layer_{lvl}_Quantity"
				unit_col = f"BOM_Layer_{lvl}_Unit"

				mat = str(row.get(mat_col, "")).strip()

				# Only real raw materials (start with 5)
				if not mat.startswith("5"):
					continue

				desc = str(row.get(desc_col, "")).strip()
				qty_str = str(row.get(qty_col, "")).strip()
				unit = str(row.get(unit_col, "")).strip().upper()

				# Skip packaging / non-RM items
				if any(k in desc.lower() for k in ("pallet", "box", "label", "paint")):
					continue

				# Parse component quantity (sum pipe-separated if present)
				try:
					qtys = [float(q) if q.strip() else 0 for q in qty_str.split(" | ")]
					component_qty = sum(qtys)
				except Exception:
					component_qty = 0

				if component_qty <= 0:
					continue

				# --- FIXED LOGIC ---
				# Layer-1 RM: use its BOM quantity as FINAL (no multiplication at all)
				if lvl == 1:
					final_qty = component_qty

				# Deeper layers: cascade using Layer-1 and intermediate layers
				else:
					cascading_multiplier = layer_1_total_qty

					# Multiply by quantities of all superior intermediate layers (2 .. lvl-1)
					for superior_lvl in range(2, lvl):
						superior_qty_col = f"BOM_Layer_{superior_lvl}_Quantity"
						superior_qty_str = str(row.get(superior_qty_col, "1")).strip()
						try:
							superior_qtys = [float(q) if q.strip() else 1 for q in superior_qty_str.split(" | ")]
							superior_total_qty = sum(superior_qtys) if superior_qtys else 1
						except Exception:
							superior_total_qty = 1
						cascading_multiplier *= superior_total_qty

					final_qty = component_qty * cascading_multiplier
				# --- END FIX ---

				if mat in rm_dict:
					rm_dict[mat]["qty"] += final_qty
				else:
					rm_dict[mat] = {
						"desc": desc,
						"qty": final_qty,
						"unit": unit,
					}

		for rm_code, details in rm_dict.items():
			if details["qty"] > 0:
				summary_records.append(
					{
						"Main_FG": main_fg,
						"Raw Material": rm_code,
						"Raw Material Description": details["desc"],
						"Raw Material qty": details["qty"],
						"Raw Material UoM": details["unit"],
					}
				)

	summary_df = pd.DataFrame(
		summary_records,
		columns=[
			"Main_FG",
			"Raw Material",
			"Raw Material Description",
			"Raw Material qty",
			"Raw Material UoM",
		],
	)

	return create_sharepoint_worksheet_optimized(
		site_id,
		file_server_relative_url,
		"FG-RM Required",
		summary_df,
		access_token,
		session_id,
		chunk_size=1000,
	)

def create_sharepoint_mb52_sheet(site_id, file_server_relative_url, session_id, access_token, stock_df):
    """Create MB52 Raw Data sheet directly in SharePoint"""
    logger.info("Creating MB52 Raw Data sheet directly in SharePoint...")
    
    return create_sharepoint_worksheet_optimized(site_id, file_server_relative_url, "MB52 Raw Data", stock_df, access_token, session_id)

def read_sharepoint_production_tracker(site_id, file_server_relative_url, session_id, access_token):
    """Read Production Tracker data robustly from SharePoint; return empty DF gracefully on error."""
    session = requests.Session()
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": ACTIVE_SESSION["id"] or session_id
    }
    range_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/usedRange"
    try:
        range_response = session.get(range_url, headers=api_headers, timeout=120)
        if range_response.status_code != 200:
            logger.error(f"Failed to read Production Tracker. Status: {range_response.status_code}. Will initialize as empty DataFrame.")
            return pd.DataFrame(), pd.DataFrame()
        range_data = range_response.json()
        values = range_data.get('values', [])
        if not values:
            logger.warning("No data found in Production Tracker")
            return pd.DataFrame(), pd.DataFrame()
        headers = values[0]
        data_rows = values[1:]
        df = pd.DataFrame(data_rows, columns=headers)
        df['Indus Part #'] = df['Indus Part #'].astype(str).str.strip().str.upper()
        if not pd.api.types.is_datetime64_any_dtype(df['DI']):
            df['DI'] = pd.to_datetime(df['DI'], errors='coerce')
        df['Status'] = df['Status'].astype(object)
        def is_blank(val):
            if pd.isna(val): return True
            if isinstance(val, str) and val.strip() == "": return True
            if isinstance(val, str) and val.strip().lower() in ("done", "cancelled"):
                return False
            return False
        mask_blank = df['Status'].apply(is_blank)
        mask_valid = (
            df['Indus Part #'].notna() &
            (df['Indus Part #'] != '') &
            (pd.to_numeric(df['Qty'], errors='coerce') > 0) &
            df['DI'].notna()
        )
        df_filtered = df[mask_blank & mask_valid].sort_values(by=['DI'], ascending=[True])
        logger.info(f"Read {len(df_filtered)} valid FIFO rows from SharePoint Production Tracker")
        return df_filtered, df
    except Exception as e:
        logger.error(f"Exception fetching Production Tracker: {e}")
        return pd.DataFrame(), pd.DataFrame()


def read_sharepoint_worksheet_data(site_id, file_server_relative_url, worksheet_name, session_id, access_token):
    """Generic function to read data from any SharePoint worksheet"""
    session = requests.Session()
    
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": ACTIVE_SESSION["id"] or session_id
    }
    
    try:
        # Get used range from specified worksheet
        range_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/usedRange"
        
        range_response = session.get(range_url, headers=api_headers, timeout=120)
        if range_response.status_code != 200:
            logger.error(f"Failed to read {worksheet_name}: {range_response.status_code}")
            return pd.DataFrame()
        
        range_data = range_response.json()
        values = range_data.get('values', [])
        
        if not values:
            logger.warning(f"No data found in {worksheet_name}")
            return pd.DataFrame()
        
        # Convert to DataFrame
        headers = values[0]
        data_rows = values[1:]
        df = pd.DataFrame(data_rows, columns=headers)
        
        logger.info(f"Read {len(df)} rows from SharePoint {worksheet_name}")
        return df
        
    except Exception as e:
        logger.error(f"Error reading SharePoint {worksheet_name}: {e}")
        return pd.DataFrame()

def update_sharepoint_production_tracker_results(site_id, file_server_relative_url, session_id, access_token, results_df, full_df):
    """Update Production Tracker columns M and N with FIFO allocation results"""
    session = requests.Session()
    
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": ACTIVE_SESSION["id"] or session_id
    }
    
    try:
        # Set headers for columns M and N
        header_updates = [
            {"range": "M1", "values": [["Material Available"]]},
            {"range": "N1", "values": [["Shortage Details"]]}
        ]
        
        for update in header_updates:
            update_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{update['range']}')"
            session.patch(update_url, headers=api_headers, json=update, timeout=60)
        
        # Create results mapping
        results_mapping = {
            int(result['index']): {
                'status': result['Status'],
                'shortage': result['Shortage_Details']
            } for _, result in results_df.iterrows()
        }
        
        # Prepare data for ALL rows
        total_rows = len(full_df)
        p_values = []
        q_values = []
        
        for idx, row in full_df.iterrows():
            status_raw = row['Status']
            status_clean = str(status_raw).strip().lower() if isinstance(status_raw, str) else ""
            
            if status_clean in ("done", "cancelled"):
                p_values.append("NA")
                q_values.append("NA")
            elif idx in results_mapping:
                p_values.append(results_mapping[idx]['status'])
                q_values.append(results_mapping[idx]['shortage'])
            else:
                p_values.append("NA")
                q_values.append("NA")
        
        # Update entire columns M and N at once
        if total_rows > 0:
            # Update column M
            p_range = f"M2:M{total_rows + 1}"
            p_values_formatted = [[val] for val in p_values]
            p_update_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{p_range}')"
            p_payload = {"values": p_values_formatted}
            
            p_response = session.patch(p_update_url, headers=api_headers, json=p_payload, timeout=300)
            if p_response.status_code in [200, 201]:
                logger.info("Column M (Material Available) updated successfully")
            
            # Update column N
            q_range = f"N2:N{total_rows + 1}"
            q_values_formatted = [[val] for val in q_values]
            q_update_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{q_range}')"
            q_payload = {"values": q_values_formatted}
            
            q_response = session.patch(q_update_url, headers=api_headers, json=q_payload, timeout=300)
            if q_response.status_code in [200, 201]:
                logger.info("Column N (Shortage Details) updated successfully")
        
        # Apply formatting to columns M and N
        apply_production_tracker_formatting(site_id, file_server_relative_url, session_id, access_token, total_rows)
        
        updated_count = len([idx for idx in range(len(full_df)) if idx in results_mapping])
        logger.info(f"Production Tracker updated with {updated_count} allocation results out of {total_rows} total rows")
        return True
        
    except Exception as e:
        logger.error(f"Error updating SharePoint Production Tracker: {e}")
        return False

def apply_production_tracker_formatting(site_id, file_server_relative_url, session_id, access_token, total_rows):
    """Apply formatting to Production Tracker columns M and N"""
    session = requests.Session()
    
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": ACTIVE_SESSION["id"] or session_id
    }
    
    try:
        # Format headers (M1:N1)
        header_range = "M1:N1"
        
        header_fill_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{header_range}')/format/fill"
        header_fill_payload = {"color": "#366092"}
        session.patch(header_fill_url, headers=api_headers, json=header_fill_payload)
        
        header_font_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{header_range}')/format/font"
        header_font_payload = {"bold": True, "color": "#FFFFFF"}
        session.patch(header_font_url, headers=api_headers, json=header_font_payload)
        
        # Data range formatting
        if total_rows > 0:
            data_range = f"M2:N{total_rows + 1}"
            
            # Apply borders
            border_ids = ["EdgeTop", "EdgeBottom", "EdgeLeft", "EdgeRight", "InsideHorizontal", "InsideVertical"]
            for border_id in border_ids:
                border_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='{data_range}')/format/borders('{border_id}')"
                border_payload = {"style": "Continuous", "color": "#000000", "weight": "Thin"}
                session.patch(border_url, headers=api_headers, json=border_payload)
            
            # Auto-fit columns M and N
            autofit_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('Production Tracker')/range(address='M:N')/format/autofitColumns"
            session.post(autofit_url, headers=api_headers)
        
        logger.info("Production Tracker columns M and N formatted successfully")
        
    except Exception as e:
        logger.warning(f"Error formatting Production Tracker columns: {e}")

def set_sharepoint_timestamp(site_id, file_server_relative_url, worksheet_name, session_id, access_token, cell_address="R1"):
    """Set timestamp in specified cell"""
    session = requests.Session()
    
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": ACTIVE_SESSION["id"] or session_id
    }
    
    update_datetime = datetime.now().strftime('Last Updated: %d-%m-%Y %H:%M:%S')
    payload = {"values": [[update_datetime]]}
    
    update_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{cell_address}')"
    
    try:
        r = session.patch(update_url, headers=api_headers, json=payload, timeout=60)
        if r.status_code in [200, 201]:
            logger.info(f"Timestamp written to cell {cell_address} in '{worksheet_name}'")
            return True
        else:
            logger.error(f"Failed to update timestamp: {r.status_code}")
            return False
    except Exception as e:
        logger.error(f"Error setting timestamp: {e}")
        return False

# --- BOM Item Data Fetch (SAME AS ORIGINAL) ---
def fetch_bom_data(api_url, username, password, verify_ssl=True, preferred_format='json', batch_size=500, max_workers=15) -> pd.DataFrame:
    logger.info("Starting BOM ITEM data fetch with enhanced concurrency...")
    session = requests.Session()
    session.auth = (username, password)
    session.verify = verify_ssl
    session.headers.update({'Accept': f'application/{preferred_format}', 'Content-Type': f'application/{preferred_format}'})
    
    select_fields = [
        'Material', 'ComponentDescription', 'BillOfMaterialComponent',
        'IsAssembly', 'BillOfMaterialItemQuantity', 'BillOfMaterialItemUnit',
        'BOMItemDescription', 'MaterialGroup'
    ]
    
    endpoint = f"{api_url}/MaterialBOMItem"
    
    def fetch_batch(skip):
        params = {'$select': ','.join(select_fields), '$top': batch_size, '$orderby': 'Material', '$skip': skip}
        try:
            logger.info(f"Fetching BOM items batch starting at {skip}...")
            response = session.get(endpoint, params=params, timeout=180)
            response.raise_for_status()
            
            content_type = response.headers.get('content-type', '')
            if 'xml' in content_type or preferred_format == 'xml':
                return parse_xml_bom(response.content)
            else:
                return parse_json_bom(response.text)
        except Exception as e:
            logger.error(f"Failed to fetch batch at skip={skip} : {e}")
            return []
    
    # Estimate total items
    params_count = {'$top': 1}
    resp_count = session.get(endpoint, params=params_count, timeout=30)
    resp_count.raise_for_status()
    data_count = resp_count.json()
    total_items = int(data_count.get('d', {}).get('__count', '0')) if '__count' in data_count.get('d', {}) else 0
    if total_items == 0:
        total_items = 50000
    
    total_batches = (total_items // batch_size) + 1
    all_records = []
    
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {executor.submit(fetch_batch, batch_num * batch_size): batch_num for batch_num in range(total_batches)}
        for future in as_completed(futures):
            records = future.result()
            if not records:
                logger.warning(f"No records fetched for batch {futures[future]}")
            else:
                all_records.extend(records)
                logger.info(f"Fetched batch size: {len(records)} (total so far: {len(all_records)})")
    
    logger.info(f"Total BOM ITEM records fetched: {len(all_records)}")
    return pd.DataFrame(all_records)

def fetch_material_bom_headers(api_url, username, password, verify_ssl=True, preferred_format='json', batch_size=500):
    logger.info("Starting BOM HEADER data fetch...")
    session = requests.Session()
    session.auth = (username, password)
    session.verify = verify_ssl
    session.headers.update({'Accept': f'application/{preferred_format}', 'Content-Type': f'application/{preferred_format}'})
    
    select_fields = ['Material', 'ProductDescription', 'BOMHeaderQuantityInBaseUnit']
    endpoint = f"{api_url}/MaterialBOM"
    params = {'$select': ','.join(select_fields), '$top': batch_size}
    
    all_records = []
    skip = 0
    
    while True:
        params['$skip'] = skip
        logger.info(f"Fetching BOM headers batch starting at {skip}...")
        response = session.get(endpoint, params=params, timeout=180)
        response.raise_for_status()
        
        content_type = response.headers.get('content-type', '')
        if 'xml' in content_type or preferred_format == 'xml':
            ns = {'atom': 'http://www.w3.org/2005/Atom', 'd': 'http://schemas.microsoft.com/ado/2007/08/dataservices', 'm': 'http://schemas.microsoft.com/ado/2007/08/dataservices/metadata'}
            tree = ET.fromstring(response.content)
            entries = tree.findall('.//atom:entry', ns)
            records = []
            for entry in entries:
                properties = entry.find('.//m:properties', ns)
                if properties is not None:
                    m = properties.find('.//d:Material', ns)
                    d = properties.find('.//d:ProductDescription', ns)
                    h = properties.find('.//d:BOMHeaderQuantityInBaseUnit', ns)
                    rec = {
                        'Material': m.text.strip() if m is not None and m.text else '',
                        'ProductDescription': d.text.strip() if d is not None and d.text else '',
                        'BOMHeaderQuantityInBaseUnit': h.text.strip() if h is not None and h.text else ''
                    }
                    records.append(rec)
        else:
            import json
            dct = response.json()
            records = [{'Material': i.get('Material', ''), 'ProductDescription': i.get('ProductDescription', ''), 'BOMHeaderQuantityInBaseUnit': i.get('BOMHeaderQuantityInBaseUnit', '')}
                       for i in dct.get('d', {}).get('results', [])]
        
        if not records:
            logger.warning("No more BOM header records found.")
            break
        
        all_records.extend(records)
        logger.info(f"Fetched header batch size: {len(records)} (total so far: {len(all_records)})")
        
        if len(records) < batch_size:
            break
        skip += batch_size
    
    logger.info(f"Total BOM HEADER records fetched: {len(all_records)}")
    return pd.DataFrame(all_records)

def fetch_material_stock_data(api_url, username, password, verify_ssl=True, preferred_format='json', batch_size=500):
    logger.info("Starting Material Stock data fetch...")
    session = requests.Session()
    session.auth = (username, password)
    session.verify = verify_ssl
    session.headers.update({'Accept': f'application/{preferred_format}', 'Content-Type': f'application/{preferred_format}'})
    
    select_fields = [
        'Material', 'Plant', 'StorageLocation', 'Batch',
        'InventorySpecialStockType',
        'InventoryStockType', 'MaterialBaseUnit', 'MatlWrhsStkQtyInMatlBaseUnit'
    ]
    
    stock_api_url = "https://my409486-api.s4hana.cloud.sap/sap/opu/odata/sap/API_MATERIAL_STOCK_SRV"
    endpoint = f"{stock_api_url}/A_MatlStkInAcctMod"
    params = {'$select': ','.join(select_fields), '$top': batch_size}
    
    all_records = []
    skip = 0
    
    while True:
        params['$skip'] = skip
        logger.info(f"Fetching Material Stock batch starting at {skip}...")
        
        try:
            response = session.get(endpoint, params=params, timeout=180)
            response.raise_for_status()
            
            content_type = response.headers.get('content-type', '')
            if 'xml' in content_type or preferred_format == 'xml':
                records = parse_xml_material_stock(response.content)
            else:
                records = parse_json_material_stock(response.text)
            
            if not records:
                logger.warning("No more Material Stock records found.")
                break
            
            all_records.extend(records)
            logger.info(f"Fetched Material Stock batch size: {len(records)} (total so far: {len(all_records)})")
            
            if len(records) < batch_size:
                break
            skip += batch_size
        except Exception as e:
            logger.error(f"Failed to fetch Material Stock batch at skip={skip}: {e}")
            break
    
    logger.info(f"Total Material Stock records fetched: {len(all_records)}")
    return pd.DataFrame(all_records)

def parse_json_bom(text):
    import json
    try:
        data = json.loads(text)
        results = []
        for item in data.get('d', {}).get('results', []):
            r = {}
            for key in ['Material', 'BillOfMaterialComponent', 'ComponentDescription', 'IsAssembly', 'BillOfMaterialItemQuantity', 'BillOfMaterialItemUnit', 'BOMItemDescription', 'MaterialGroup']:
                value = item.get(key, '')
                if isinstance(value, dict): value = str(value)
                r[key] = value
            results.append(r)
        return results
    except Exception as ex:
        logger.error(f"JSON BOM parse failed: {str(ex)}")
        return []

def parse_xml_bom(xml_content):
    ns = {'atom': 'http://www.w3.org/2005/Atom', 'd': 'http://schemas.microsoft.com/ado/2007/08/dataservices', 'm': 'http://schemas.microsoft.com/ado/2007/08/dataservices/metadata'}
    tree = ET.fromstring(xml_content)
    records = []
    entries = tree.findall('.//atom:entry', ns)
    
    for entry in entries:
        properties = entry.find('.//m:properties', ns)
        if properties is not None:
            record = {}
            for field in ['Material', 'BillOfMaterialComponent', 'ComponentDescription', 'IsAssembly', 'BillOfMaterialItemQuantity', 'BillOfMaterialItemUnit', 'BOMItemDescription', 'MaterialGroup']:
                elem = properties.find(f'.//d:{field}', ns)
                record[field] = elem.text.strip() if elem is not None and elem.text else ''
            records.append(record)
    return records

def parse_json_material_stock(text):
    import json
    try:
        data = json.loads(text)
        results = []
        for item in data.get('d', {}).get('results', []):
            r = {}
            for key in ['Material', 'Plant', 'StorageLocation', 'Batch', 'InventorySpecialStockType', 'InventoryStockType', 'MaterialBaseUnit', 'MatlWrhsStkQtyInMatlBaseUnit']:
                value = item.get(key, '')
                if isinstance(value, dict):
                    value = str(value)
                r[key] = value
            results.append(r)
        return results
    except Exception as ex:
        logger.error(f"JSON Material Stock parse failed: {str(ex)}")
        return []

def parse_xml_material_stock(xml_content):
    ns = {'atom': 'http://www.w3.org/2005/Atom', 'd': 'http://schemas.microsoft.com/ado/2007/08/dataservices', 'm': 'http://schemas.microsoft.com/ado/2007/08/dataservices/metadata'}
    tree = ET.fromstring(xml_content)
    records = []
    entries = tree.findall('.//atom:entry', ns)
    
    for entry in entries:
        properties = entry.find('.//m:properties', ns)
        if properties is not None:
            record = {}
            for field in ['Material', 'Plant', 'StorageLocation', 'Batch', 'InventorySpecialStockType', 'InventoryStockType', 'MaterialBaseUnit', 'MatlWrhsStkQtyInMatlBaseUnit']:
                elem = properties.find(f'.//d:{field}', ns)
                record[field] = elem.text.strip() if elem is not None and elem.text else ''
            records.append(record)
    return records

def standardize_bom_data(df):
    logger.info("Standardizing BOM item dataframe...")
    df = df.fillna('')
    
    for col in ['Material', 'BillOfMaterialComponent', 'ComponentDescription', 'BillOfMaterialItemUnit']:
        df[col] = df[col].astype(str).str.strip().str.upper()
    
    if 'BOMItemDescription' not in df.columns:
        df['BOMItemDescription'] = ''
    df['BOMItemDescription'] = df['BOMItemDescription'].astype(str).str.strip()

    if 'MaterialGroup' not in df.columns:
        df['MaterialGroup'] = ''
    df['MaterialGroup'] = df['MaterialGroup'].astype(str).str.strip().str.upper()
    # a component's material group is the same everywhere, so fill blanks from any line that has it
    _mg_known = df[df['MaterialGroup'] != ''].groupby('BillOfMaterialComponent')['MaterialGroup'].first()
    df['MaterialGroup'] = df['MaterialGroup'].where(df['MaterialGroup'] != '', df['BillOfMaterialComponent'].map(_mg_known).fillna(''))
    _mg_filled = df.drop_duplicates('BillOfMaterialComponent')
    logger.info(f"Material group filled for {(_mg_filled['MaterialGroup'] != '').sum()} of {len(_mg_filled)} components; "
                f"{(_mg_filled['MaterialGroup'].map(material_group_desc) != '').sum()} found in the description list")

    df['BillOfMaterialItemQuantity'] = pd.to_numeric(df['BillOfMaterialItemQuantity'], errors='coerce').fillna(0)
    df['IsAssembly'] = df['IsAssembly'].astype(str).str.upper()
    df = df.drop_duplicates(subset=['Material', 'BillOfMaterialComponent'], keep='first')
    
    logger.info(f"BOM item dataframe standardized, total rows: {len(df)}")
    return df

def tree_explode(df, start_mat, max_depth, parents_seen=None, path=None):
    if parents_seen is None:
        parents_seen = set()
    if path is None:
        path = []
    
    children = df[df["Material"] == start_mat]
    
    if len(children) == 0:
        yield path
    else:
        for idx, row in children.iterrows():
            child = row["BillOfMaterialComponent"]
            if child == "" or child in parents_seen:
                continue
            
            row_tuple = (
                start_mat,
                row["ComponentDescription"],
                row["BillOfMaterialItemQuantity"],
                row["BillOfMaterialItemUnit"],
                child,
                row["IsAssembly"],
                row["BOMItemDescription"],
                row["MaterialGroup"]
            )
            
            next_path = path + [row_tuple]
            
            if row["IsAssembly"] == 'X' and len(next_path) < max_depth:
                yield from tree_explode(df, child, max_depth, parents_seen | {start_mat}, next_path)
            else:
                yield next_path

def create_mb52_stock_dict(mb52_df):
    """Create stock dictionary from MB52 data for Plant 1000 ONLY"""
    logger.info("Creating MB52 stock dictionary for Plant 1000...")
    
    plant_1000_stock = mb52_df[
        mb52_df['Plant'].astype(str).str.strip() == '1000'
    ].copy()
    
    logger.info(f"Plant 1000 stock records: {len(plant_1000_stock)}")
    
    qty_col = 'MatlWrhsStkQtyInMatlBaseUnit'
    if qty_col in plant_1000_stock.columns:
        plant_1000_stock[qty_col] = pd.to_numeric(plant_1000_stock[qty_col], errors='coerce').fillna(0)
    else:
        logger.error(f"Column {qty_col} not found in stock data")
        return {}
    
    # Aggregate by Material
    stock_dict = {}
    for _, row in plant_1000_stock.iterrows():
        material = str(row['Material']).strip().upper()
        qty = float(row[qty_col])
        
        if material in stock_dict:
            stock_dict[material] += qty
        else:
            stock_dict[material] = qty
    
    logger.info(f"MB52 stock dictionary created: {len(stock_dict)} unique materials")
    return stock_dict

def get_rm_requirements_for_fg(fg, fg_rm_df):
    """Get RM requirements for a given FG"""
    rm_req = {}
    fg_norm = str(fg).strip().upper()
    
    rows = fg_rm_df[fg_rm_df['Main_FG'].astype(str).str.strip().str.upper() == fg_norm]
    
    for _, row in rows.iterrows():
        rm_code = str(row['Raw Material']).strip()
        rm_desc = str(row['Raw Material Description']).strip()
        rm_qty = float(row['Raw Material qty']) if pd.notna(row['Raw Material qty']) else 0.0
        rm_uom = str(row['Raw Material UoM']).strip()
        
        if rm_code and rm_qty > 0:
            rm_req[rm_code] = (rm_desc, rm_qty, rm_uom)
    
    return rm_req

def set_fixed_header_timestamp(site_id, file_server_relative_url, worksheet_name, cell_address, session_id, access_token):
    """
    Writes timestamp ONLY to the given header cell (e.g., 'R1') and applies header formatting.
    Usage:
      set_fixed_header_timestamp(SITE_ID, SHAREPOINT_FILE, "Production Tracker", "R1", session_id, access_token)
    """

    session = requests.Session()
    api_headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "workbook-session-id": ACTIVE_SESSION["id"] or session_id
    }

    # 1) Write timestamp value
    ts_text = "Last Updated: " + datetime.now().strftime('%d-%m-%Y %H:%M:%S')
    put_url = (
        f"https://graph.microsoft.com/v1.0/sites/{site_id}"
        f"/drive/root:/{file_server_relative_url}:/workbook/worksheets('{worksheet_name}')/range(address='{cell_address}')"
    )
    put_resp = session.patch(put_url, headers=api_headers, json={"values": [[ts_text]]}, timeout=60)
    if put_resp.status_code not in (200, 201):
        logger.error(f"Timestamp write failed for {worksheet_name} {cell_address}: {put_resp.status_code} {put_resp.text}")
        return False

    # 2) Apply header formatting (fill + font) to that exact cell
    fill_url = f"{put_url}/format/fill"
    font_url = f"{put_url}/format/font"

    # Header background color
    session.patch(fill_url, headers=api_headers, json={"color": "#366092"}, timeout=30)
    # Header font style
    session.patch(font_url, headers=api_headers, json={"bold": True, "color": "#FFFFFF"}, timeout=30)

    logger.info(f"Timestamp written and formatted at {worksheet_name} {cell_address}")
    return True


def process_fifo_allocation(df_orders, fg_rm_df, mb52_df):
    """FIFO ALLOCATION: Date-based FIFO across ALL orders"""
    logger.info("=== STARTING SHAREPOINT-BASED FIFO ALLOCATION PROCESSING ===")
    
    # Create initial MB52 stock dictionary for Plant 1000
    mb52_stock_dict = create_mb52_stock_dict(mb52_df)
    initial_stock_total = sum(mb52_stock_dict.values())
    
    # Initialize tracking variables
    results = []
    allocation_summary = {
        'total_orders': len(df_orders),
        'successful_orders': 0,
        'failed_orders': 0,
        'by_material': {},
        'stock_consumed': 0
    }
    
    logger.info(f"Processing {len(df_orders)} orders in strict DI date order (FIFO)")
    logger.info(f"Initial total stock available: {initial_stock_total:.3f} units across all materials")
    
    # Track statistics by material
    material_stats = {}
    
    # Process each order in DI date sequence
    for order_sequence, (original_idx, order) in enumerate(df_orders.iterrows(), 1):
        indus_part = str(order['Indus Part #']).strip().upper()
        order_qty = float(order['Qty'])
        order_di = order['DI']
        
        logger.info(f"\n--- Processing Order {order_sequence}/{len(df_orders)}: {indus_part} ---")
        logger.info(f"Order Details: Qty={order_qty}, DI={order_di}, Original Index={original_idx}")
        
        if indus_part not in material_stats:
            material_stats[indus_part] = {'successful': 0, 'total': 0}
        material_stats[indus_part]['total'] += 1
        
        # Get RM requirements for this FG
        rm_requirements = get_rm_requirements_for_fg(indus_part, fg_rm_df)
        
        if not rm_requirements:
            logger.warning(f"No RM requirements found for {indus_part} - marking as Error")
            results.append({
                'index': original_idx,
                'Indus Part #': indus_part,
                'Qty': order_qty,
                'DI': order_di,
                'Status': 'Error',
                'Shortage_Details': f'No RM requirements found in FG-RM Required sheet for {indus_part}'
            })
            continue
        
        logger.info(f"Found {len(rm_requirements)} RM components for {indus_part}")
        
        # Apply UOM filtering
        rm_requirements_filtered = {}
        excluded_count = 0
        
        for rm_code, (rm_desc, qty_per_fg, rm_uom) in rm_requirements.items():
            unit = str(rm_uom).strip().upper()
            if unit in ("L", "LT", "LITRE", "ROL", "ROLL"):
                excluded_count += 1
                logger.debug(f"Excluding {rm_code} due to UoM: {unit}")
                continue
            rm_requirements_filtered[rm_code] = (rm_desc, qty_per_fg, rm_uom)
        
        if excluded_count > 0:
            logger.info(f"Excluded {excluded_count} components due to UoM filtering")
        
        if not rm_requirements_filtered:
            logger.warning(f"No eligible RM components for {indus_part} after UoM filtering - marking as Error")
            results.append({
                'index': original_idx,
                'Indus Part #': indus_part,
                'Qty': order_qty,
                'DI': order_di,
                'Status': 'Error',
                'Shortage_Details': f'No eligible RM components (excluded UoMs: L/ROL) for {indus_part}'
            })
            continue
        
        logger.info(f"Processing {len(rm_requirements_filtered)} eligible RM components")
        
        # Calculate total RM requirements for this order
        can_fulfill_completely = True
        shortage_components = []
        required_materials = {}
        total_stock_needed = 0
        
        for rm_code, (rm_desc, qty_per_fg, rm_uom) in rm_requirements_filtered.items():
            total_required = qty_per_fg * order_qty
            available_stock = mb52_stock_dict.get(rm_code, 0)
            total_stock_needed += total_required
            
            required_materials[rm_code] = {
                'required': total_required,
                'available': available_stock,
                'desc': rm_desc,
                'uom': rm_uom
            }
            
            logger.debug(f" {rm_code}: Required={total_required:.3f} {rm_uom}, Available={available_stock:.3f}")
            
            if available_stock < total_required:
                can_fulfill_completely = False
                shortage_qty = total_required - available_stock
                shortage_components.append(f"{rm_desc} ({rm_code}): {format_shortage_quantity(shortage_qty, rm_uom)} {rm_uom}")
                logger.info(f" SHORTAGE: {rm_code} needs {total_required:.3f}, only {available_stock:.3f} available")
        
        logger.info(f"Total stock needed for order: {total_stock_needed:.3f} units")
        
        # CRITICAL FIFO LOGIC: Reserve stock regardless of shortage status
        stock_consumed_this_order = 0
        for rm_code, details in required_materials.items():
            stock_to_consume = min(details['required'], details['available'])
            if stock_to_consume > 0:
                mb52_stock_dict[rm_code] -= stock_to_consume
                stock_consumed_this_order += stock_to_consume
                logger.debug(f" Reserved {stock_to_consume:.3f} of {rm_code}, remaining: {mb52_stock_dict[rm_code]:.3f}")
        
        allocation_summary['stock_consumed'] += stock_consumed_this_order
        
        # Determine order status and record results
        if can_fulfill_completely:
            logger.info(f"Order FULFILLED: All RM components available, consumed {stock_consumed_this_order:.3f} units")
            results.append({
                'index': original_idx,
                'Indus Part #': indus_part,
                'Qty': order_qty,
                'DI': order_di,
                'Status': 'YES',
                'Shortage_Details': ''
            })
            allocation_summary['successful_orders'] += 1
            material_stats[indus_part]['successful'] += 1
        else:
            shortage_text = ", ".join(shortage_components)
            logger.info(f"Order FAILED: Component shortages detected, consumed {stock_consumed_this_order:.3f} units")
            logger.info(f"Shortage details: {shortage_text}")
            results.append({
                'index': original_idx,
                'Indus Part #': indus_part,
                'Qty': order_qty,
                'DI': order_di,
                'Status': 'NO',
                'Shortage_Details': shortage_text
            })
            allocation_summary['failed_orders'] += 1
        
        # Progress logging
        if order_sequence % 10 == 0 or order_sequence == len(df_orders):
            remaining_stock_total = sum(mb52_stock_dict.values())
            stock_utilization = ((initial_stock_total - remaining_stock_total) / initial_stock_total * 100) if initial_stock_total > 0 else 0
            logger.info(f"Progress: {order_sequence}/{len(df_orders)} orders processed, {stock_utilization:.1f}% stock utilized")
    
    # Compile final results
    allocation_summary['by_material'] = material_stats
    
    # Convert results to DataFrame and sort by original index
    results_df = pd.DataFrame(results)
    if not results_df.empty:
        results_df = results_df.sort_values('index').reset_index(drop=True)
    
    # Final summary logging
    final_stock_total = sum(mb52_stock_dict.values())
    total_consumed = initial_stock_total - final_stock_total
    stock_efficiency = (total_consumed / initial_stock_total * 100) if initial_stock_total > 0 else 0
    
    logger.info("\n=== SHAREPOINT-BASED FIFO ALLOCATION COMPLETED ===")
    logger.info(f"Stock Utilization Summary:")
    logger.info(f" Initial Stock: {initial_stock_total:.3f} units")
    logger.info(f" Consumed Stock: {total_consumed:.3f} units ({stock_efficiency:.1f}%)")
    logger.info(f" Remaining Stock: {final_stock_total:.3f} units")
    logger.info(f"Order Fulfillment Summary:")
    logger.info(f" Total Orders: {allocation_summary['total_orders']}")
    logger.info(f" Successful: {allocation_summary['successful_orders']} ({allocation_summary['successful_orders']/allocation_summary['total_orders']*100:.1f}%)")
    logger.info(f" Failed: {allocation_summary['failed_orders']} ({allocation_summary['failed_orders']/allocation_summary['total_orders']*100:.1f}%)")
    
    return results_df, allocation_summary

def format_shortage_quantity(qty, uom):
    """Return shortage quantity rounded per UOM: PC=1, KG=2, else=2"""
    try:
        qty = float(qty)
        if str(uom).strip().upper() == 'PC':
            return f"{qty:.1f}"
        elif str(uom).strip().upper() == 'KG':
            return f"{qty:.2f}"
        else:
            return f"{qty:.2f}"
    except Exception:
        return str(qty)

def validate_data_linkage(df_tracker, fg_rm_df):
    """Enhanced data validation with detailed logging"""
    logger.info("=== VALIDATING DATA LINKAGE ===")
    
    try:
        tracker_parts = set(df_tracker['Indus Part #'].astype(str).str.strip().str.upper().unique())
        fg_rm_parts = set(fg_rm_df['Main_FG'].astype(str).str.strip().str.upper().unique())
        
        logger.info(f"Production Tracker unique parts: {len(tracker_parts)}")
        logger.info(f"FG-RM Required unique parts: {len(fg_rm_parts)}")
        
        # Find matches and mismatches
        matches = tracker_parts.intersection(fg_rm_parts)
        tracker_only = tracker_parts - fg_rm_parts
        fg_rm_only = fg_rm_parts - tracker_parts
        
        logger.info(f"Matching parts: {len(matches)}")
        logger.info(f"Tracker-only parts: {len(tracker_only)}")
        logger.info(f"FG-RM-only parts: {len(fg_rm_only)}")
        
        if matches:
            logger.info(f"Sample matches: {list(matches)[:5]}")
        if tracker_only:
            logger.warning(f"Parts in Tracker but not in FG-RM Required: {list(tracker_only)[:5]}")
        
        return len(matches) > 0, matches
        
    except Exception as e:
        logger.error(f"Data validation error: {e}")
        return False, set()

# =============================================================================
# PART SUMMARY (one row per part number) + ASSEMBLY PARTS (one row per coil inside a combi/assembly)
# Built in memory from the "BOM Layer-wise" data, then written to the same workbook.
# =============================================================================
_COIL_KIND_RX = re.compile(r"^\s*(EVAPORATOR|CONDENSER|HEATER)\s+COIL", re.I)
_TUBE_RX = re.compile(r"\b(CU|AL)\s?(IG|PL)\s*([\d\.]+)\s*X\s*([\d\.]+)", re.I)
_CTR_RX = re.compile(r"CTR[-\s]*([\d\.]+)|([\d\.]+)\s*CTR", re.I)
_FOIL_RX = re.compile(r"FOIL\s+(\d+(?:\.\d+)?)\s*X\s*(\d+(?:\.\d+)?)\s*(.*)", re.I)


def _ps_float(v):
    try:
        x = float(v)
        return 0.0 if x != x else x
    except Exception:
        return 0.0


def _ps_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v != v:
        return ""
    return str(v).strip()


def _ps_kind(desc, mg=""):
    """Coarse kind of one BOM line, used only by the summary sheets."""
    d = _ps_str(desc).upper()
    if _COIL_KIND_RX.search(d):
        return "coil"
    if is_foil_component(d, mg):
        return "foil"
    if re.match(r"^\s*(HPS\d?|TPHPS)\b", d):
        return "hps"      # hairpin straight
    if re.match(r"^\s*(HPC\d?|TPHPC)\b", d):
        return "hpc"      # hairpin cross
    if re.match(r"^\s*STR\d?\b", d):
        return "straight"
    if re.search(r"U[ -]?BEND", d):
        return "ubend"
    if "ENDPLATE" in d:
        return "endplate"
    if re.match(r"^\s*(INLET|INTLET|ASSY INLET)", d):
        return "inlet"
    if re.match(r"^\s*(OUTLET|OUTLT|ASSY OUTLET)", d):
        return "outlet"
    if "PALLET" in d:
        return "pallet"
    if re.match(r"^\s*(HCP\s+)?BOX\b", d) or "CARTON" in d:
        return "box"
    return ""


def _ps_round(x):
    return None if x is None else round(float(x), 4)


def _ps_qty_text(q):
    return str(int(q)) if float(q).is_integer() else str(round(float(q), 3))


def _ps_size_text(w, t):
    """'395 x 0.11' from width and thickness in mm (blank if either is missing)."""
    if w is None or t is None:
        return None
    return f"{w:g} x {t:g}"


def _ps_metrics(lines):
    """Metrics for the direct components of one node (the part itself, or one coil inside an assembly)."""
    m = {"foil_code": "", "foil_w": None, "foil_t": None, "foil_coat": "", "foil_kg": None,
         "fpi": None, "fins": None, "len": None, "surface": "", "edge": "", "check": "",
         "hps": 0.0, "hpc": 0.0, "hairpin_detail": "", "hps_spec": "", "hpc_spec": "",
         "straight": 0.0, "ubends": 0.0,
         "tube_mat": "", "tube_type": "", "tube_od": None, "tube_wall": None, "tube_ctr": None, "tube_spec": "",
         "endplate": "", "inlet": "No", "outlet": "No", "box": "", "pallet": "", "notes": []}
    foils = [l for l in lines if l["kind"] == "foil"]
    if foils:
        foils.sort(key=lambda l: abs(l["qty"]), reverse=True)
        f = foils[0]
        m["foil_code"] = f["code"]
        m["foil_kg"] = _ps_round(sum(x["qty"] for x in foils if x["code"] == f["code"]))
        mm = _FOIL_RX.search(f["desc"])
        if mm:
            # 'AL FOIL 395 X 0.11 HYDROPHILLIC' -> roll width 395 mm, thickness 0.11 mm
            m["foil_w"], m["foil_t"] = float(mm.group(1)), float(mm.group(2))
            rest = mm.group(3).upper()
            m["foil_coat"] = ("Hydrophilic" if "HYDROPHIL" in rest else "Hydrophobic" if "HYDROPHOB" in rest
                              else "Bare" if "BARE" in rest else "")
        fin = parse_fin_spec(f["itemtext"])
        m["fpi"], m["fins"], m["len"] = fin["fpi"], fin["fins"], fin["length_mm"]
        m["surface"], m["edge"], m["check"] = fin["surface"], fin["edge"], fin["check"]
        if len({x["code"] for x in foils}) > 1:
            m["notes"].append("More than one foil code: largest quantity shown")
    detail, tin_plated = {}, False
    for l in lines:
        k = l["kind"]
        if k == "hps":
            m["hps"] += l["qty"]
        elif k == "hpc":
            m["hpc"] += l["qty"]
        elif k == "straight":
            m["straight"] += l["qty"]
        elif k == "ubend":
            m["ubends"] += l["qty"]
        if k in ("hps", "hpc"):
            tok = re.match(r"^\s*([A-Z0-9]+)", l["desc"].upper())
            tok = tok.group(1) if tok else k.upper()
            detail[tok] = detail.get(tok, 0.0) + l["qty"]
            if tok.startswith("TP"):
                tin_plated = True
    # every hairpin variant on its own (HPS, HPS1, HPS2, HPC, TPHPS ...), straight family first
    m["hairpin_detail"] = " | ".join(
        f"{t} x{_ps_qty_text(q)}"
        for t, q in sorted(detail.items(), key=lambda x: (0 if "HPS" in x[0] else 1, x[0])))
    if tin_plated:
        m["notes"].append("Includes tin-plated hairpins (TPHPS/TPHPC), counted with HPS/HPC")
    # tube text of each hairpin type, as typed on the BOM
    for kind, key in (("hps", "hps_spec"), ("hpc", "hpc_spec")):
        for l in lines:
            if l["kind"] == kind and l["itemtext"]:
                m[key] = l["itemtext"]
                break
    # parsed tube size: from the straight-hairpin line, else the cross-hairpin line, else the straight tube
    for kind in ("hps", "hpc", "straight"):
        for l in lines:
            if l["kind"] == kind and l["itemtext"]:
                m["tube_spec"] = l["itemtext"]
                mt = _TUBE_RX.search(l["itemtext"])
                if mt:
                    m["tube_mat"] = "Copper" if mt.group(1).upper() == "CU" else "Aluminium"
                    m["tube_type"] = "Inner-grooved" if mt.group(2).upper() == "IG" else "Plain"
                    m["tube_od"], m["tube_wall"] = float(mt.group(3)), float(mt.group(4))
                    mc = _CTR_RX.search(l["itemtext"])
                    if mc:
                        m["tube_ctr"] = float(mc.group(1) or mc.group(2))
                break
        if m["tube_spec"]:
            break
    eps = []
    for l in lines:
        if l["kind"] == "endplate":
            mat = re.match(r"^\s*(GI|AL|ALUMINI?UM|SS|CU)\b", l["desc"].upper())
            if mat:
                eps.append({"AL": "AL", "ALUMINIUM": "AL", "ALUMINUM": "AL"}.get(mat.group(1), mat.group(1)))
    m["endplate"] = " / ".join(sorted(set(eps)))
    m["inlet"] = "Yes" if any(l["kind"] == "inlet" for l in lines) else "No"
    m["outlet"] = "Yes" if any(l["kind"] == "outlet" for l in lines) else "No"
    m["box"] = " | ".join(dict.fromkeys(l["desc"] for l in lines if l["kind"] == "box"))
    m["pallet"] = " | ".join(dict.fromkeys(l["desc"] for l in lines if l["kind"] == "pallet"))
    return m


def build_part_summary(df_layerwise):
    """From the BOM Layer-wise rows, build (Part Summary, Assembly Parts) DataFrames."""
    part_cols = ["Part No", "Description", "Part Type", "Coils Inside", "Foil Code",
                 "Foil Size mm (width x thickness)", "Foil Width mm", "Foil Thickness mm", "Foil Coating", "Foil kg",
                 "FPI", "Fins per Pack", "Fin Length mm", "Fin Surface", "Fin Edge", "Fin Spec Check",
                 "Hairpins HPS (Straight)", "Hairpins HPC (Cross)", "Hairpin Detail",
                 "HPS Tube Spec (as typed)", "HPC Tube Spec (as typed)",
                 "Straight Tubes", "U-Bends",
                 "Tube Material", "Tube Type", "Tube OD mm", "Tube Wall mm", "Tube CTR mm", "Tube Spec (as typed)",
                 "Endplate Material", "Inlet", "Outlet", "Box", "Pallet",
                 "Total HPS (all coils)", "Total HPC (all coils)", "Total U-Bends (all coils)",
                 "Total Foil kg (all coils)", "Notes"]
    asm_cols = ["Parent Part No", "Parent Description", "Child Part No", "Child Description", "Child Type", "Qty",
                "Child In Part Summary", "Foil Code", "Foil Size mm (width x thickness)", "Foil Width mm",
                "Foil Thickness mm", "Foil Coating", "Foil kg", "FPI", "Fins per Pack", "Fin Length mm",
                "Fin Surface", "Fin Edge", "Fin Spec Check",
                "Hairpins HPS (Straight)", "Hairpins HPC (Cross)",
                "HPS Tube Spec (as typed)", "HPC Tube Spec (as typed)", "Straight Tubes", "U-Bends"]
    if df_layerwise is None or df_layerwise.empty or "Main_FG" not in df_layerwise.columns:
        return pd.DataFrame(columns=part_cols), pd.DataFrame(columns=asm_cols)

    L = range(1, MAX_LAYERS + 1)

    def col(r, name):
        return r.get(name, "")

    products, descs = {}, {}
    for r in df_layerwise.to_dict("records"):
        fg = _ps_str(r.get("Main_FG"))
        if not fg:
            continue
        descs.setdefault(fg, _ps_str(r.get("Main_FG_Description")))
        bucket = products.setdefault(fg, {"seen": set(), "lines": []})
        key = []
        for n in L:
            code = _ps_str(col(r, f"BOM_Layer_{n}"))
            if not code:
                break
            key.append(code)
            tkey = tuple(key)
            if tkey in bucket["seen"]:
                continue
            bucket["seen"].add(tkey)
            desc = _ps_str(col(r, f"BOM_Layer_{n}_Desc"))
            mg = _ps_str(col(r, f"BOM_Layer_{n}_MaterialGroup"))
            bucket["lines"].append({"layer": n, "key": tkey, "parent_key": tkey[:-1], "code": code, "desc": desc,
                                    "qty": _ps_float(col(r, f"BOM_Layer_{n}_Quantity")),
                                    "itemtext": _ps_str(col(r, f"BOM_Layer_{n}_ItemText")),
                                    "kind": _ps_kind(desc, mg)})

    part_rows, asm_rows = [], []
    known = set(products)
    for fg in sorted(products):
        lines = products[fg]["lines"]
        l1 = [l for l in lines if l["layer"] == 1]
        coils = [l for l in l1 if l["kind"] == "coil"]
        own = _ps_metrics([l for l in l1 if l["kind"] != "coil"])
        desc = descs.get(fg, "")
        notes = list(own["notes"])

        if coils:
            ptype = "Assembly of coils"
        elif re.search(r"COMBI|ASSEMBLY COIL", desc, re.I):
            ptype = "Combi built directly"
        elif re.search(r"(EVAPORATOR|CONDENSER|HEATER) COIL", desc, re.I):
            ptype = "Single coil"
        else:
            ptype = "Other"

        # fall back: foil one level down inside a non-coil sub-assembly
        if not coils and not own["foil_code"]:
            l2 = [l for l in lines if l["layer"] == 2 and l["kind"] == "foil"]
            if l2:
                alt = _ps_metrics(l2)
                for k in ("foil_code", "foil_w", "foil_t", "foil_coat", "foil_kg", "fpi", "fins", "len",
                          "surface", "edge", "check"):
                    own[k] = alt[k]
                notes.append("Foil found one level down")
        if not coils and not own["foil_code"]:
            notes.append("No foil line found")

        tot_hps, tot_hpc, tot_u, tot_f = own["hps"], own["hpc"], own["ubends"], (own["foil_kg"] or 0.0)
        for c in coils:
            sub = [l for l in lines if l["layer"] == 2 and l["parent_key"] == c["key"]]
            cm = _ps_metrics(sub)
            q = c["qty"] if c["qty"] else 1.0
            tot_hps += q * cm["hps"]
            tot_hpc += q * cm["hpc"]
            tot_u += q * cm["ubends"]
            tot_f += q * (cm["foil_kg"] or 0.0)
            ctype = (_COIL_KIND_RX.match(c["desc"]).group(1).capitalize()
                     if _COIL_KIND_RX.match(c["desc"]) else "Coil")
            if c["code"] not in known:
                notes.append(f"Coil {c['code']} has no row of its own")
            asm_rows.append([fg, desc, c["code"], c["desc"], ctype, _ps_round(c["qty"]),
                             "Yes" if c["code"] in known else "No",
                             cm["foil_code"] or None, _ps_size_text(cm["foil_w"], cm["foil_t"]),
                             cm["foil_w"], cm["foil_t"], cm["foil_coat"] or None, cm["foil_kg"],
                             cm["fpi"], cm["fins"], cm["len"], cm["surface"] or None, cm["edge"] or None,
                             cm["check"] or None,
                             _ps_round(cm["hps"]), _ps_round(cm["hpc"]),
                             cm["hps_spec"] or None, cm["hpc_spec"] or None,
                             _ps_round(cm["straight"]), _ps_round(cm["ubends"])])

        is_asm = ptype == "Assembly of coils"
        blank = (lambda v: None if is_asm else v)
        part_rows.append([
            fg, desc, ptype, len(coils) if is_asm else None,
            blank(own["foil_code"] or None),
            blank(_ps_size_text(own["foil_w"], own["foil_t"])), blank(own["foil_w"]), blank(own["foil_t"]),
            blank(own["foil_coat"] or None), blank(own["foil_kg"]),
            blank(own["fpi"]), blank(own["fins"]), blank(own["len"]),
            blank(own["surface"] or None), blank(own["edge"] or None), blank(own["check"] or None),
            blank(_ps_round(own["hps"])), blank(_ps_round(own["hpc"])), blank(own["hairpin_detail"] or None),
            blank(own["hps_spec"] or None), blank(own["hpc_spec"] or None),
            blank(_ps_round(own["straight"])), blank(_ps_round(own["ubends"])),
            blank(own["tube_mat"] or None), blank(own["tube_type"] or None), blank(own["tube_od"]),
            blank(own["tube_wall"]), blank(own["tube_ctr"]), blank(own["tube_spec"] or None),
            own["endplate"] or None, blank(own["inlet"]), blank(own["outlet"]),
            own["box"] or None, own["pallet"] or None,
            _ps_round(tot_hps), _ps_round(tot_hpc), _ps_round(tot_u), _ps_round(tot_f),
            "; ".join(dict.fromkeys(notes)) or None])

    return (pd.DataFrame(part_rows, columns=part_cols), pd.DataFrame(asm_rows, columns=asm_cols))


# --- MAIN FUNCTION: SHAREPOINT-ONLY PROCESSING ---
def main():
    """COMPLETE SHAREPOINT-ONLY BOM layer-wise FIFO allocation system"""
    logger.info("Starting SHAREPOINT-ONLY BOM layer-wise FIFO allocation system...")
    start_time = time.time()
    
    # SharePoint configuration
    SITE_ID = "201c767f-23cc-4eff-b76c-47a275706531,094b21f9-6ed4-47c3-a751-3b217a18b98e"
    SHAREPOINT_FILE = "Manufacturing Plan New.xlsx"
    
    try:
        # === PHASE 1: DATA COLLECTION ===
        logger.info("=== DATA COLLECTION PHASE ===")
        
        # Enhanced concurrent data fetching
        df_bom = fetch_bom_data(SAP_API_URL, USERNAME, PASSWORD, VERIFY_SSL, PREFERRED_FORMAT, BATCH_SIZE, max_workers=15)
        df_bom = standardize_bom_data(df_bom)
        
        # Fetch other data concurrently
        with ThreadPoolExecutor(max_workers=2) as executor:
            header_future = executor.submit(fetch_material_bom_headers, SAP_API_URL, USERNAME, PASSWORD, VERIFY_SSL, PREFERRED_FORMAT, BATCH_SIZE)
            stock_future = executor.submit(fetch_material_stock_data, SAP_API_URL, USERNAME, PASSWORD, VERIFY_SSL, PREFERRED_FORMAT, BATCH_SIZE)
            
            header_df = header_future.result()
            stock_df = stock_future.result()
        
        header_mapping = dict(zip(header_df["Material"], header_df["ProductDescription"]))
        header_qty_mapping = dict(zip(header_df["Material"], pd.to_numeric(header_df["BOMHeaderQuantityInBaseUnit"], errors="coerce").fillna(1)))
        
        logger.info(f"Data collection completed - BOM: {len(df_bom)}, Headers: {len(header_df)}, Stock: {len(stock_df)}")
        
        # === PHASE 2: BOM TREE STRUCTURE CREATION ===
        logger.info("=== BOM TREE STRUCTURE CREATION ===")
        
        # Create BOM tree structure (in memory, no local files)
        logger.info("Creating BOM layer-wise structure...")
        all_paths = []
        for fg in df_bom["Material"].unique():
            fg_paths = list(tree_explode(df_bom, fg, MAX_LAYERS))
            all_paths.extend(fg_paths)
        
        logger.info(f"BOM tree structure created with {len(all_paths)} paths")
        
        # === PHASE 3: SHAREPOINT EXCEL SESSION CREATION ===
        logger.info("=== SHAREPOINT EXCEL SESSION CREATION ===")
        
        # Get access token and create session
        access_token = get_access_token()
        session_id = None
        
        # Create SharePoint Excel session
        for session_attempt in range(2):
            try:
                session_url = f"https://graph.microsoft.com/v1.0/sites/{SITE_ID}/drive/root:/{SHAREPOINT_FILE}:/workbook/createSession"
                session_headers = {
                    "Authorization": f"Bearer {access_token}",
                    "Content-Type": "application/json"
                }
                session_payload = {"persistChanges": True}
                
                session_response = requests.post(session_url, headers=session_headers, json=session_payload, timeout=30)
                if session_response.status_code in [200, 201]:
                    session_id = session_response.json().get("id")
                    logger.info(f"Created SharePoint Excel session: {session_id}")
                    break
                else:
                    if session_attempt < 1:
                        time.sleep(2)
                        continue
                    else:
                        logger.error("Failed to create SharePoint session")
                        return
            except Exception as session_error:
                if session_attempt < 1:
                    time.sleep(2)
                    continue
                else:
                    logger.error(f"Failed to create SharePoint session: {session_error}")
                    return
        
        if not session_id:
            logger.error("Could not create SharePoint Excel session")
            return
            
        try:
            # === PHASE 4: SHAREPOINT SHEET CREATION ===
            logger.info("=== SHAREPOINT SHEET CREATION ===")
            
            # Create BOM Layer-wise sheet directly in SharePoint
            logger.info("Creating BOM Layer-wise sheet in SharePoint...")
            success_main = create_sharepoint_bom_layerwise_sheet(SITE_ID, SHAREPOINT_FILE, session_id, access_token, df_bom, all_paths, MAX_LAYERS, header_mapping, header_qty_mapping)
            time.sleep(2)
            
            # Read the created BOM Layer-wise data for FG-RM processing
            logger.info("Reading created BOM Layer-wise data...")
            df_main_sheet = read_sharepoint_worksheet_data(SITE_ID, SHAREPOINT_FILE, "BOM Layer-wise", session_id, access_token)
            if df_main_sheet.empty or "Main_FG" not in df_main_sheet.columns:
                logger.error("BOM Layer-wise could not be read back, so the run stops here and the other sheets are left unchanged. "
                             "Close the workbook everywhere and run again.")
                return
            
            # Create FG-RM Required sheet
            logger.info("Creating FG-RM Required sheet in SharePoint...")
            success_fg_rm = create_sharepoint_fg_rm_sheet(SITE_ID, SHAREPOINT_FILE, session_id, access_token, df_main_sheet, MAX_LAYERS)
            time.sleep(2)
            
            # Create MB52 Raw Data sheet
            logger.info("Creating MB52 Raw Data sheet in SharePoint...")
            success_stock = create_sharepoint_mb52_sheet(SITE_ID, SHAREPOINT_FILE, session_id, access_token, stock_df)
            time.sleep(1)

            # Create Part Summary + Assembly Parts sheets (a failure here never stops the rest of the run)
            try:
                if df_main_sheet.empty:
                    logger.warning("BOM Layer-wise data not available - Part Summary skipped")
                else:
                    logger.info("Creating Part Summary and Assembly Parts sheets in SharePoint...")
                    part_df, assembly_df = build_part_summary(df_main_sheet)
                    logger.info(f"Part Summary rows: {len(part_df)}, Assembly Parts rows: {len(assembly_df)}")
                    ok_part = create_sharepoint_worksheet_optimized(SITE_ID, SHAREPOINT_FILE, "Part Summary", part_df, access_token, session_id, native_types=True)
                    set_fixed_header_timestamp(SITE_ID, SHAREPOINT_FILE, "Part Summary", f"{get_excel_column_letter(len(part_df.columns) + 2)}1", session_id, access_token)
                    ok_asm = create_sharepoint_worksheet_optimized(SITE_ID, SHAREPOINT_FILE, "Assembly Parts", assembly_df, access_token, session_id, native_types=True)
                    set_fixed_header_timestamp(SITE_ID, SHAREPOINT_FILE, "Assembly Parts", f"{get_excel_column_letter(len(assembly_df.columns) + 2)}1", session_id, access_token)
                    logger.info(f"Part Summary written: {ok_part}; Assembly Parts written: {ok_asm}")
            except Exception as summary_error:
                logger.error(f"Part Summary step failed (rest of the run continues): {summary_error}")
            
            # === PHASE 5: SHAREPOINT DATA VALIDATION ===
            logger.info("=== SHAREPOINT DATA VALIDATION ===")
            
            # Read Production Tracker data from SharePoint
            df_orders, full_df = read_sharepoint_production_tracker(SITE_ID, SHAREPOINT_FILE, session_id, access_token)
            
            if df_orders.empty:
                logger.error("No SharePoint Production Tracker data available for processing")
                return
            
            logger.info(f"Ready to process {len(df_orders)} orders from SharePoint Production Tracker")
            
            # Read FG-RM Required and MB52 data from SharePoint
            fg_rm_df = read_sharepoint_worksheet_data(SITE_ID, SHAREPOINT_FILE, "FG-RM Required", session_id, access_token)
            mb52_df = read_sharepoint_worksheet_data(SITE_ID, SHAREPOINT_FILE, "MB52 Raw Data", session_id, access_token)
            
            # Validate data linkage
            is_valid, matches = validate_data_linkage(df_orders, fg_rm_df)
            if not is_valid:
                logger.error("CRITICAL: No data linkage between Production Tracker and FG-RM Required!")
                return
            
            logger.info(f"Data validation passed: {len(matches)} matching materials")
            
            # === PHASE 6: SHAREPOINT FIFO ALLOCATION EXECUTION ===
            logger.info("=== SHAREPOINT FIFO ALLOCATION EXECUTION ===")
            
            # Execute the core FIFO allocation logic using SharePoint data
            results_df, allocation_summary = process_fifo_allocation(df_orders, fg_rm_df, mb52_df)
            
            # === PHASE 7: SHAREPOINT RESULTS UPDATE ===
            logger.info("=== SHAREPOINT RESULTS UPDATE ===")
            
            # Update SharePoint Production Tracker with results
            success_tracker = update_sharepoint_production_tracker_results(SITE_ID, SHAREPOINT_FILE, session_id, access_token, results_df, full_df)
            
            # After finishing Production Tracker updates
            set_fixed_header_timestamp(SITE_ID, SHAREPOINT_FILE, "Production Tracker", "O1", session_id, access_token)

            # After finishing BOM Layer-wise updates
            set_fixed_header_timestamp(SITE_ID, SHAREPOINT_FILE, "BOM Layer-wise", "AL1", session_id, access_token)

            # After finishing FG-RM Required updates
            set_fixed_header_timestamp(SITE_ID, SHAREPOINT_FILE, "FG-RM Required", "F1", session_id, access_token)

            # After finishing MB52 Raw Data updates
            set_fixed_header_timestamp(SITE_ID, SHAREPOINT_FILE, "MB52 Raw Data", "I1", session_id, access_token)


            
            # Log SharePoint operations results
            if success_main and success_fg_rm and success_stock and success_tracker:
                logger.info("All SharePoint Excel operations completed successfully")
            else:
                logger.warning("Some SharePoint Excel operations failed")
                
        finally:
            # Always close the session
            try:
                close_session_url = f"https://graph.microsoft.com/v1.0/sites/{SITE_ID}/drive/root:/{SHAREPOINT_FILE}:/workbook/closeSession"
                close_headers = {
                    "Authorization": f"Bearer {access_token}",
                    "workbook-session-id": ACTIVE_SESSION["id"] or session_id
                }
                close_response = requests.post(close_session_url, headers=close_headers, timeout=15)
                if close_response.status_code in [200, 204]:
                    logger.info("SharePoint Excel session closed successfully")
            except Exception as close_error:
                logger.warning(f"Error closing session: {close_error}")
        
        # === PHASE 8: FINAL SUMMARY REPORT ===
        logger.info("=== SHAREPOINT-ONLY ALLOCATION SUMMARY REPORT ===")
        logger.info(f" Total Orders: {allocation_summary['total_orders']}")
        logger.info(f" Successful: {allocation_summary['successful_orders']} ({allocation_summary['successful_orders']/allocation_summary['total_orders']*100:.1f}%)")
        logger.info(f" Failed: {allocation_summary['failed_orders']} ({allocation_summary['failed_orders']/allocation_summary['total_orders']*100:.1f}%)")
        
        if not results_df.empty:
            total_qty_requested = results_df['Qty'].sum()
            successful_qty = results_df[results_df['Status'] == 'YES']['Qty'].sum()
            failed_qty = results_df[results_df['Status'] == 'NO']['Qty'].sum()
            logger.info(f" Total Quantity Requested: {total_qty_requested}")
            logger.info(f" Quantity Allocated: {successful_qty}")
            logger.info(f" Quantity Short: {failed_qty}")
        
        elapsed_time = time.time() - start_time
        logger.info(f"SHAREPOINT-ONLY FIFO allocation system completed successfully in {elapsed_time:.2f} seconds")
        
    except Exception as e:
        logger.error(f"Critical error in main process: {e}")
        import traceback
        logger.error(f"Full traceback: {traceback.format_exc()}")
        raise

if __name__ == "__main__":
    main()
